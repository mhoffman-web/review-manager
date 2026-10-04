"""Mention detection, template suggestion, saved views, archive, employee report, AI draft (mocked)."""
from datetime import datetime, timedelta

import pytest
from sqlalchemy import select

from app import auth
from app.db import SessionLocal, engine, init_db
from app.models import (AiRule, Base, Employee, Location, ReplyTemplate, Response, Review, ReviewMention,
                        ReviewSourceLink, SavedView, SiteGroup, User)
from app.daterange import resolve_range
from app.reports import build_trends, employee_report, monthly_summary, rating_distribution, responder_stats, window_report
from app.text_intel import apply_intel, build_roster, classify_theme, detect_mentions, suggest_templates


@pytest.fixture()
def db():
    Base.metadata.drop_all(engine)
    init_db()
    s = SessionLocal()
    loc = Location(name="ICON Thompson Lane", brand="ICON", state="TN", city="Nashville")
    s.add(loc); s.flush()
    link = ReviewSourceLink(location_id=loc.id, source="google", external_account_id="1", external_location_id="L1", display_name="Icon Car Wash", active=True)
    s.add(link)
    s.add(Employee(name="Jared", location_id=loc.id))
    s.add(Employee(name="Gregory Banks", aliases="Mr. Banks;Greg", location_id=loc.id))
    s.add(User(email="lily@x.com", name="Lily Collins", password_hash=auth.hash_password("dev-password-lily-2026"), role="admin"))
    s.commit()
    yield s, loc, link
    s.close()


def mk(s, link, ext, rating, text, author="Pat Q", days_ago=1.0, replied=False):
    t = datetime.utcnow() - timedelta(days=days_ago)
    r = Review(source_link_id=link.id, source="google", external_id=ext, author_name=author, rating=rating, text=text,
               created_at_source=t, updated_at_source=t, has_owner_reply=replied,
               owner_reply_text="thanks" if replied else None, owner_reply_updated_at=t + timedelta(hours=20) if replied else None, raw_json="{}")
    s.add(r); s.flush()
    return r


def test_detect_mentions_cues_and_roster():
    roster = {"jared": "Jared", "gregory banks": "Gregory Banks", "mr. banks": "Gregory Banks", "greg": "Gregory Banks"}
    assert detect_mentions("Jared is the man fast Service and great attitude", "Eric Crain", roster) == [("Jared", "Jared")]
    assert detect_mentions("Today I had great help from Mr. Banks, he explained everything.", "Chin Suk Kim", roster) == [("Gregory Banks", "Gregory Banks")]
    found = detect_mentions("shout out to Karla for being informative on the great deals", "Isabel Benitez", roster)
    assert found == [("Karla", None)]
    # author's own name and ordinary words are not mentions
    assert detect_mentions("Great wash great staff. Isabel was here Tuesday.", "Isabel Benitez", roster) == []
    assert detect_mentions("The user didn't write a review", None, roster) == []
    assert detect_mentions(None, None, roster) == []
    # two names
    names = [n for n, _ in detect_mentions("Lucien and Damon were super helpful as well.", "Michael McClure", {})]
    assert "Lucien" in names and "Damon" in names


def test_classify_theme_and_suggestions(db):
    s, loc, link = db
    assert classify_theme("Waited 25 minutes in line, one lane open") == "Long Line"
    assert classify_theme("Charged twice and nobody answered the phone") == "Billing/Cancellation"
    assert classify_theme("Great wash") == "Unknown"
    assert classify_theme(None) == "No Content"
    assert classify_theme("Dryer left water streaks everywhere") == "Dryer"
    assert classify_theme("Kiosk screen would not take my card") == "POS"
    tpls = [
        ReplyTemplate(name="5 general", min_rating=5, max_rating=5, tags="general", body="Thanks {first_name}"),
        ReplyTemplate(name="5 employee", min_rating=5, max_rating=5, tags="employee", body="Thanks for recognising {employee}"),
        ReplyTemplate(name="5 rating only", min_rating=5, max_rating=5, tags="no_comment;general", body="Thanks for the 5 stars {first_name}"),
        ReplyTemplate(name="Sorry wait", min_rating=1, max_rating=3, tags="wait", body="Sorry about the wait at {site}"),
        ReplyTemplate(name="Sorry billing", min_rating=1, max_rating=3, tags="billing", body="Sorry about billing"),
    ]
    for t in tpls:
        s.add(t)
    s.commit()
    r = mk(s, link, "a", 5, "Jared is the man fast Service and great attitude", author="Eric Crain")
    assert apply_intel(s, r, build_roster(s)) == 1
    s.refresh(r)
    ranked = suggest_templates(r, tpls, r.mention_names)
    assert ranked[0][0].name == "5 employee"
    assert ranked[0][0].render(r, "Lily") == "Thanks for recognising Jared"
    r2 = mk(s, link, "b", 5, None, author="Derek")
    assert suggest_templates(r2, tpls, [])[0][0].name == "5 rating only"
    r3 = mk(s, link, "c", 1, "Charged twice this month. Called and nobody answered.", author="Sam K")
    apply_intel(s, r3, build_roster(s))
    assert r3.category == "Billing/Cancellation"
    assert suggest_templates(r3, tpls, [])[0][0].name == "Sorry billing"


def test_employee_report_and_monthly(db):
    s, loc, link = db
    roster = build_roster(s)
    for i, text in enumerate(["Jared was awesome", "Thanks to Jared and Eli for the help", "Greg explained the plans well", "Great wash", None]):
        r = mk(s, link, f"m{i}", 5, text, author="Casey L", days_ago=2 + i, replied=True)
        apply_intel(s, r, roster)
    s.commit()
    rep = employee_report(s, resolve_range("last30"))
    names = {row["name"]: row for row in rep["rows"]}
    assert names["Jared"]["count"] == 2 and names["Jared"]["on_roster"]
    assert names["Gregory Banks"]["count"] == 1
    assert names["Eli"]["on_roster"] is False
    assert rep["reviews_with_mentions"] == 3 and rep["reviews_with_text"] == 4
    months = monthly_summary(s, months=3)
    assert months and months[0]["n"] >= 1 and months[0]["response_rate"] == 100
    w = window_report(s, resolve_range("last30"))
    assert w.count == 5 and w.response_rate == 100 and w.mention_reviews == 3
    t = build_trends(s, resolve_range("last30"))
    assert t["granularity"] == "day" and sum(t["counts"]["ICON"]) == 5 and len(t["labels"]) == 30
    assert build_trends(s, resolve_range("last_quarter"))["granularity"] == "week"
    # a custom window is whole LOCAL days: fixed dates around both edges, never "now"-relative
    from zoneinfo import ZoneInfo
    from app.config import settings
    tz, utc = ZoneInfo(settings.timezone), ZoneInfo("UTC")
    for ext, local in [("e1", datetime(2026, 6, 9, 23, 30)), ("e2", datetime(2026, 6, 10, 0, 30)),
                       ("e3", datetime(2026, 6, 12, 23, 30)), ("e4", datetime(2026, 6, 13, 0, 30))]:
        t = local.replace(tzinfo=tz).astimezone(utc).replace(tzinfo=None)
        s.add(Review(source_link_id=link.id, source="google", external_id=ext, rating=4, text="edge",
                     created_at_source=t, updated_at_source=t, raw_json="{}"))
    s.commit()
    w2 = window_report(s, resolve_range("custom", "2026-06-10", "2026-06-12"))
    assert w2.count == 2


def test_web_views_archive_draft_and_admin(db, monkeypatch):
    from fastapi.testclient import TestClient
    from app.web import app
    import app.web as web
    s, loc, link = db
    u = s.execute(select(User)).scalar_one()
    r = mk(s, link, "w1", 1, "Line wrapped around the building", author="Taylor R")
    r_ok = mk(s, link, "w2", 5, None, author="Derek", days_ago=0.1)
    s.add(ReplyTemplate(name="Sorry wait", min_rating=1, max_rating=3, tags="wait", body="Sorry about the wait at {site}"))
    s.add(AiRule(text="Keep it under 25 words."))
    s.commit()
    c = TestClient(app)
    c.cookies.set(auth.COOKIE_NAME, auth.make_session_cookie(u))
    # saved view round trip
    resp = c.post("/views", data={"name": "LW Negatives", "is_shared": "1", "params": '{"view":"negative","days":7}'}, follow_redirects=False)
    assert resp.status_code == 303 and "sv=" in resp.headers["location"]
    page = c.get(resp.headers["location"])
    assert page.status_code == 200 and "LW Negatives" in page.text and "Line wrapped" in page.text
    # archive removes from unanswered and from counts
    before = c.get("/?view=unanswered").text.count("Rating only, no written review")
    c.post(f"/reviews/{r_ok.id}/archive", data={"back": ""}, follow_redirects=False)
    assert c.get("/?view=unanswered").text.count("Rating only, no written review") == before - 1
    assert "Archived" in c.get(f"/reviews/{r_ok.id}").text
    # draft endpoint: disabled without a key, works with a mocked model
    assert c.post(f"/reviews/{r.id}/draft").status_code == 503
    monkeypatch.setattr(web, "draft_reply", lambda review, rules, examples, agent: "We're sorry about the line, Taylor. Call (615) 776-7837.")
    monkeypatch.setattr(web.settings, "anthropic_api_key", "test-key")
    j = c.post(f"/reviews/{r.id}/draft").json()
    assert j["ok"] and "Taylor" in j["text"]
    # template usage increments on post (fake adapter)
    monkeypatch.setattr(web, "get_adapter", lambda src: type("A", (), {"post_reply": lambda self, l, e, t: datetime.utcnow()})())
    tpl = s.execute(select(ReplyTemplate)).scalar_one()
    c.post(f"/reviews/{r.id}/reply", data={"text": "Sorry about the wait", "template_id": str(tpl.id), "ai_generated": "0"}, follow_redirects=False)
    s.expire_all()
    assert s.get(ReplyTemplate, tpl.id).usage_count == 1
    assert responder_stats(s, resolve_range("today"))[0]["count"] == 1
    # admin pages + employee promote + group
    for path in ["/admin/employees", "/admin/groups", "/admin/ai", "/admin/templates", "/reports", "/reports/employees", f"/sites/{loc.id}"]:
        assert c.get(path).status_code == 200, path
    c.post("/admin/employees", data={"name": "Eli", "location_id": str(loc.id), "aliases": "", "role": "", "active": "1"}, follow_redirects=False)
    # quick group creation from the inbox (any signed-in user)
    r2 = c.post("/groups", data={"name": "Nashville", "location_ids": [str(loc.id)], "back": "/?view=all"}, follow_redirects=False)
    assert r2.status_code == 303 and "group_id=" in r2.headers["location"]
    s.expire_all()
    assert s.execute(select(Employee).where(Employee.name == "Eli")).scalar_one_or_none() is not None
    g = s.execute(select(SiteGroup)).scalar_one()
    assert [l.id for l in g.locations] == [loc.id]
    assert c.get(f"/?group_id={g.id}&view=all").status_code == 200
    # date-range filters on inbox and reports, including a custom window
    assert "Line wrapped" in c.get("/?view=all&range=last7").text
    assert "Line wrapped" not in c.get("/?view=all&range=custom&start=2020-01-01&end=2020-01-31").text
    for path in ["/reports?range=last_month", "/reports?range=custom&start=2026-09-01&end=2026-09-15&brand=ICON", f"/sites/{loc.id}?range=ytd",
                 "/reports/employees?range=last_week", f"/reports/distribution?range=last30&location_id={loc.id}&group_id={g.id}", "/reports/morning?edition=tn",
                 "/reports/morning?edition=il&range=last7", "/admin/recipients"]:
        assert c.get(path).status_code == 200, path
    # multi-select inbox filters: two ratings, a site and a group at once
    multi = c.get(f"/?view=all&rating=1&rating=5&location_id={loc.id}&group_id={g.id}&brand=ICON").text
    assert "Line wrapped" in multi and "Nice wash" not in multi
    assert "NPS" not in c.get("/reports").text
    dist = rating_distribution(s, resolve_range("last30"), location_ids=[loc.id])
    assert dist["rows"][0]["name"] == "ICON Thompson Lane" and sum(dist["totals"]) == dist["total"] and dist["chart"]["labels"] == ["ICON Thompson Lane"]
    r3 = c.post("/admin/recipients", data={"email": "tn@x.com", "name": "TN", "edition": "tn"}, follow_redirects=False)
    assert r3.status_code == 303
    s.expire_all()
    from app.models import ReportRecipient
    assert s.execute(select(ReportRecipient).where(ReportRecipient.email == "tn@x.com")).scalar_one().edition == "tn"
    # per-edition paste box: several addresses at once, "Name <email>" keeps the name, one person on two editions
    r4 = c.post("/admin/recipients", data={"emails": "Lily <lily@x.com>, sarah@x.com\nnot-an-email", "edition": "il"}, follow_redirects=False)
    assert r4.status_code == 303 and "added=2" in r4.headers["location"] and "skipped=not-an-email" in r4.headers["location"]
    c.post("/admin/recipients", data={"emails": "LILY@x.com", "edition": "tn"}, follow_redirects=False)
    s.expire_all()
    lily = s.execute(select(ReportRecipient).where(ReportRecipient.email == "lily@x.com")).scalar_one()
    assert lily.name == "Lily" and lily.editions == ["il", "tn"]
    from app.reports import recipient_groups
    groups = recipient_groups(s)
    assert groups["il"] == ["lily@x.com", "sarah@x.com"] and set(groups["tn"]) == {"lily@x.com", "tn@x.com"}
    page = c.get("/admin/recipients").text
    assert page.count("lily@x.com") == 2 and "Add to IL" in page and "Add to Corporate" in page
    assert c.post(f"/admin/recipients/{lily.id}/delete", data={"edition": "il"}, follow_redirects=False).status_code == 303
    s.expire_all()
    assert lily.editions == ["tn"]
    c.post(f"/admin/recipients/{lily.id}/delete", follow_redirects=False)
    s.expire_all()
    assert s.execute(select(ReportRecipient).where(ReportRecipient.email == "lily@x.com")).scalar_one_or_none() is None
    mk(s, link, "w3", 4, "Nice wash, quick line", author="Kim Z", days_ago=0.2); s.commit()
    page = c.get("/?view=all").text
    assert "Reviewer" in page and 'data-reply="' in page and "Claim" not in page and "New site group" in page
    # ---- round three: queue-aware prev/next, inline JSON reply, guardrails, events, removed view, report, exports, API
    from app.models import ReviewEvent, ApiKey
    from app.api import new_key, hash_key
    w3 = s.execute(select(Review).where(Review.external_id == "w3")).scalar_one()
    ctx = "view=all&q=&range=&start=&end=&sort=&dir=asc&sv=0"
    detail = c.get(f"/reviews/{w3.id}?ctx={ctx}").text
    assert "in this list" in detail and "Activity" in detail and "Report to Google" in detail
    # guardrails: an unfilled placeholder and a wrong greeting come back as warnings (422 on the JSON path) until force=1
    j = c.post(f"/reviews/{w3.id}/reply", data={"text": "Hi Taylor, thanks! {employee} was glad to help."}, headers={"Accept": "application/json"})
    assert j.status_code == 422 and any("placeholder" in w for w in j.json()["warnings"]) and any("greets Taylor" in w for w in j.json()["warnings"])
    j = c.post(f"/reviews/{w3.id}/reply", data={"text": "Hi Kim, thanks for the kind words!", "force": "1"}, headers={"Accept": "application/json"})
    assert j.status_code == 200 and j.json()["ok"] and j.json()["reply"]["by"] == "Lily"
    s.expire_all()
    w3 = s.get(Review, w3.id)
    assert w3.has_owner_reply and w3.first_replied_at is not None
    first_at = w3.first_replied_at
    # editing keeps the first-reply time (response time does not move) and logs an edit event
    c.post(f"/reviews/{w3.id}/reply", data={"text": "Hi Kim, thanks so much for the kind words!", "force": "1"}, follow_redirects=False)
    s.expire_all(); w3 = s.get(Review, w3.id)
    kinds = [e.kind for e in w3.events]
    assert w3.first_replied_at == first_at and "reply_posted" in kinds and "reply_edited" in kinds
    # archive toggles with an undo link and an event; JSON form used by the keyboard shortcut
    a = c.post(f"/reviews/{r.id}/archive", headers={"Accept": "application/json"}).json()
    assert a["ok"] and a["archived"] and isinstance(a["undo_archive"], int)
    c.post(f"/reviews/{a['undo_archive']}/archive", headers={"Accept": "application/json"})
    s.expire_all()
    assert s.get(Review, r.id).is_archived is False and [e.kind for e in s.get(Review, r.id).events][-2:] == ["archived", "unarchived"]
    # report to Google: disputed reviews leave the figures until Google decides
    before_total = rating_distribution(s, resolve_range("last30"), location_ids=[loc.id])["total"]
    c.post(f"/reviews/{r.id}/report", data={"action": "reported", "note": "not a customer"}, follow_redirects=False)
    s.rollback()          # end this session's read transaction so it sees the web request's write
    assert s.get(Review, r.id).is_disputed and rating_distribution(s, resolve_range("last30"), location_ids=[loc.id])["total"] == before_total - 1
    assert "Disputed" in c.get("/?view=all").text
    c.post(f"/reviews/{r.id}/report", data={"action": "kept"}, follow_redirects=False)
    s.rollback()
    assert s.get(Review, r.id).report_status == "kept" and rating_distribution(s, resolve_range("last30"), location_ids=[loc.id])["total"] == before_total
    # a hand-set theme survives re-classification and feeds the AI examples
    c.post(f"/reviews/{r.id}/note", data={"note": "", "category": "Dryer"}, follow_redirects=False)
    s.expire_all()
    assert s.get(Review, r.id).category_source == "manual"
    from app.sync import _manual_examples
    assert ("Line wrapped around the building", "Dryer") in _manual_examples(s)
    # removed view, exports, weekly, alerts preview, admin pages all render
    r.is_deleted, r.removed_at = True, datetime.utcnow(); s.commit()
    assert "Removed" in c.get("/").text and "No longer on Google" in c.get("/?view=removed").text
    for path in ["/reports/weekly", f"/reports/weekly?as_of=2026-09-30&location_id={loc.id}", "/reports/alerts/preview", "/admin/sites", "/admin/api",
                 "/export/reviews.csv?view=all", "/export/distribution.csv?range=last30", "/locations"]:
        assert c.get(path).status_code == 200, path
    x = c.get("/export/reviews.xlsx?view=all")
    assert x.status_code == 200 and x.headers["content-type"].startswith("application/vnd.openxmlformats") and len(x.content) > 2000
    assert "Line wrapped" in c.get("/export/reviews.csv?view=removed").text
    # admin: add a site and record a listing by id
    c.post("/admin/sites", data={"name": "ICON Murfreesboro", "brand": "ICON", "state": "tn", "city": "Murfreesboro", "snowflake_location_ids": "", "active": "1"}, follow_redirects=False)
    s.expire_all()
    new_loc = s.execute(select(Location).where(Location.name == "ICON Murfreesboro")).scalar_one()
    assert new_loc.state == "TN"
    c.post("/admin/sites/listing", data={"source": "google", "external_location_id": "locations/99887766", "display_name": "ICON Car Wash Murfreesboro", "location_id": str(new_loc.id)}, follow_redirects=False)
    s.expire_all()
    nl = s.execute(select(ReviewSourceLink).where(ReviewSourceLink.external_location_id == "99887766")).scalar_one()
    assert nl.active and nl.location_id == new_loc.id
    # a Wash N' Roll title is recorded but never mapped
    c.post("/admin/sites/listing", data={"source": "google", "external_location_id": "5551212", "display_name": "Wash N' Roll Car Wash - Goodlettsville", "location_id": str(new_loc.id)}, follow_redirects=False)
    s.expire_all()
    wnr = s.execute(select(ReviewSourceLink).where(ReviewSourceLink.external_location_id == "5551212")).scalar_one()
    assert not wnr.active and wnr.location_id is None
    # read-only API with a key
    raw = new_key()
    s.add(ApiKey(name="Teams Portal", prefix=raw[:11], key_hash=hash_key(raw))); s.commit()
    assert c.get("/api/v1/summary").status_code == 401 and c.get("/api/v1/summary", headers={"X-API-Key": "rm_nope"}).status_code == 403
    summ = c.get("/api/v1/summary?range=last30", headers={"X-API-Key": raw}).json()
    assert summ["totals"]["reviews"] >= 1 and len(summ["totals"]["distribution"]) == 5
    assert any(x["name"] == "ICON Thompson Lane" and x["reviews"] >= 1 for x in summ["sites"])
    rv = c.get("/api/v1/reviews?range=last30&limit=5", headers={"Authorization": f"Bearer {raw}"}).json()
    assert rv["count"] >= 1 and {"site", "rating", "text", "replied", "link"} <= set(rv["reviews"][0])
    assert c.get("/api/v1/sites", headers={"X-API-Key": raw}).json()["sites"][0]["brand"] == "ICON"
    assert c.get("/api/v1/leaderboard?range=last30", headers={"X-API-Key": raw}).status_code == 200


def test_inbox_server_sort(db):
    from fastapi.testclient import TestClient
    from app.web import app
    s, loc, link = db
    u = s.execute(select(User)).scalar_one()
    mk(s, link, "s1", 5, "Alpha great", author="Zed A", days_ago=3)
    mk(s, link, "s2", 1, "Beta awful", author="Amy B", days_ago=1)
    mk(s, link, "s3", 3, "Gamma meh", author="Mel C", days_ago=2)
    s.commit()
    c = TestClient(app)
    c.cookies.set(auth.COOKIE_NAME, auth.make_session_cookie(u))
    html = c.get("/?view=all&sort=rating&dir=asc").text
    assert html.index("Beta awful") < html.index("Gamma meh") < html.index("Alpha great")
    html = c.get("/?view=all&sort=rating&dir=desc").text
    assert html.index("Alpha great") < html.index("Gamma meh") < html.index("Beta awful")
    html = c.get("/?view=all&sort=author&dir=asc").text
    assert html.index("Amy B") < html.index("Mel C") < html.index("Zed A")
    assert "sorted by reviewer" in html and 'data-server-sort' in html
    assert c.get("/?view=all&sort=bogus").status_code == 200


def test_redirect_targets_and_script_json_are_safe(db):
    """Open-redirect and script-breakout regressions."""
    from fastapi.testclient import TestClient
    from app.web import app, _safe_path
    for bad in ("//evil.com", "/\\evil.com", "https://evil.com", "javascript:alert(1)", "evil.com", "/\tx", ""):
        assert _safe_path(bad) == "/", bad
    assert _safe_path("/?view=open&q=a%26b") == "/?view=open&q=a%26b"
    s, loc, link = db
    rid = mk(s, link, "xss", 2, "dirty car").id; s.commit()
    c = TestClient(app)
    r = c.post("/login", data={"email": "lily@x.com", "password": "dev-password-lily-2026", "next": "//evil.com"}, follow_redirects=False)
    assert r.headers["location"] == "/"
    r = c.post(f"/reviews/{rid}/claim", data={"back": "//evil.com/x"}, follow_redirects=False)
    assert r.headers["location"] == f"/reviews/{rid}"
    r = c.post(f"/reviews/{rid}/archive", data={"back": "https://evil.com"}, follow_redirects=False)
    assert r.headers["location"].startswith(f"/reviews/{rid}?") and "undo_archive=" in r.headers["location"]
    assert "undo=%2F" not in r.headers["location"]
    # a reviewer name that tries to close the script block is escaped inside the page's JSON
    rv = s.get(Review, rid)
    rv.author_name = "</script><script>alert(1)</script>"; s.commit()
    page = c.get(f"/reviews/{rid}").text
    assert "</script><script>alert(1)" not in page


def test_api_keys_never_travel_in_urls(db):
    from fastapi.testclient import TestClient
    from app.web import app
    s, loc, link = db
    c = TestClient(app)
    c.post("/login", data={"email": "lily@x.com", "password": "dev-password-lily-2026"})
    r = c.post("/admin/api", data={"name": "Teams Portal"}, follow_redirects=False)
    assert r.status_code == 200 and r.headers.get("cache-control") == "no-store"
    import re
    raw = re.search(r'id="newKey"[^>]*>([^<]+)<', r.text).group(1)
    assert raw.startswith("rm_") and raw not in c.get("/admin/api").text
    anon = TestClient(app)
    assert anon.get(f"/api/v1/sites?key={raw}").status_code == 400
    assert anon.get("/api/v1/sites", headers={"X-API-Key": raw}).status_code == 200


def test_search_values_survive_paging_and_post_and_next(db, monkeypatch):
    """A search containing & # + % must stay intact in sort links, exports and the review queue."""
    from fastapi.testclient import TestClient
    from urllib.parse import parse_qs, urlsplit
    from app.web import app
    import app.web as web
    monkeypatch.setattr(web, "get_adapter", lambda src: type("A", (), {"post_reply": lambda self, l, e, t: datetime.utcnow()})())
    import html as _html, re
    s, loc, link = db
    a = mk(s, link, "q1", 2, "A&W #1 50% + more dirt", days_ago=2)
    b = mk(s, link, "q2", 1, "A&W #1 50% + more dirt again", days_ago=1)
    mk(s, link, "q3", 1, "unrelated", days_ago=0.5)
    s.commit()
    c = TestClient(app)
    c.post("/login", data={"email": "lily@x.com", "password": "dev-password-lily-2026"})
    page = c.get("/", params={"view": "all", "q": "A&W #1 50% +"}).text
    assert "q=A%26W+%231+50%25+%2B" in page
    href = _html.unescape(re.search(r'href="(/reviews/%d\?ctx=[^"]+)"' % b.id, page).group(1))
    ctx = parse_qs(urlsplit(href).query)["ctx"][0]
    assert parse_qs(ctx)["q"] == ["A&W #1 50% +"]
    rp = c.get(href).text
    assert 'id="navNext"' in rp and f"/reviews/{a.id}?ctx=" in rp          # the filtered queue, not the whole inbox
    r = c.post(f"/reviews/{b.id}/reply", data={"text": "Thanks for letting us know, we are on it.", "go_next": "1", "ctx": ctx, "force": "1"},
               follow_redirects=False)
    loc_hdr = r.headers["location"]
    assert loc_hdr.startswith(f"/reviews/{a.id}?")
    assert parse_qs(parse_qs(urlsplit(loc_hdr).query)["ctx"][0])["q"] == ["A&W #1 50% +"]


def test_exports_never_contain_live_formulas(db):
    import io
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from app.web import app
    s, loc, link = db
    mk(s, link, "f1", 1, '=HYPERLINK("http://evil.example","click")', author="=cmd|' /C calc'!A0", days_ago=1)
    mk(s, link, "f2", 2, "-2+3 bad wash", author="@SUM(1)", days_ago=1)
    s.commit()
    c = TestClient(app)
    c.post("/login", data={"email": "lily@x.com", "password": "dev-password-lily-2026"})
    body = c.get("/export/reviews.csv?view=all").text
    assert "'=HYPERLINK" in body and "'=cmd" in body and "'@SUM(1)" in body and "'-2+3 bad wash" in body
    assert ',=' not in body and '"=' not in body
    wb = load_workbook(io.BytesIO(c.get("/export/reviews.xlsx?view=all").content))
    cells = [cell for row in wb.active.iter_rows() for cell in row]
    hyper = [cl for cl in cells if isinstance(cl.value, str) and cl.value.startswith("=HYPERLINK")]
    assert hyper and all(cl.data_type == "s" for cl in hyper)
    assert not any(cl.data_type == "f" for cl in cells)


def test_bulk_classify_respects_manual_themes(db, monkeypatch):
    from app import ai
    from app.sync import classify_window
    s, loc, link = db
    manual = mk(s, link, "c1", 1, "the dryer left water everywhere"); manual.category, manual.category_source = "Customer Service", "manual"
    auto = mk(s, link, "c2", 2, "line took forty minutes")
    empty = mk(s, link, "c3", 1, None)
    s.commit()
    monkeypatch.setattr(ai, "classify_negative", lambda r, cats, examples=None: "Long Line")
    out = classify_window(s, [manual, auto, empty], force=True)
    assert out == {"classified": 2, "kept": 0, "manual": 1, "failed": 0}
    assert manual.category == "Customer Service" and manual.category_source == "manual"
    assert (auto.category, auto.category_source) == ("Long Line", "ai")
    assert (empty.category, empty.category_source) == ("No Content", "keyword")


def test_ai_cut_off_answers_are_unavailable_not_unknown(db, monkeypatch):
    import types, sys
    from app import ai
    s, loc, link = db
    r = mk(s, link, "t1", 1, "awful"); s.commit()
    calls = {}

    class Msgs:
        def create(self, **kw):
            calls.update(kw)
            return types.SimpleNamespace(stop_reason="max_tokens", content=[])
    fake = types.SimpleNamespace(Anthropic=lambda api_key: types.SimpleNamespace(messages=Msgs()),
                                 RateLimitError=type("R", (Exception,), {}), APIStatusError=type("S", (Exception,), {}),
                                 APIConnectionError=type("C", (Exception,), {}))
    monkeypatch.setitem(sys.modules, "anthropic", fake)
    monkeypatch.setattr(ai.settings, "anthropic_api_key", "test")
    with pytest.raises(ai.AiUnavailable):
        ai.classify_negative(r, ["Long Line", "Unknown"])
    assert calls["max_tokens"] >= 1000 and calls["output_config"] == {"effort": "low"}
    with pytest.raises(ai.AiUnavailable):
        ai.draft_reply(r)
    assert calls["max_tokens"] >= 2000
