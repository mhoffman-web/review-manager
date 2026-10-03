"""FastAPI web UI: inbox, review detail + reply, reports, sites, admin."""
from __future__ import annotations

import csv
import io
import json
import logging
import re
from collections import Counter
from datetime import datetime, timedelta
from urllib.parse import parse_qs, urlencode
from pathlib import Path
from typing import Dict, List, Optional
from zoneinfo import ZoneInfo

from fastapi import BackgroundTasks, Depends, FastAPI, Form, HTTPException, Query, Request, status
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, Response as RawResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, selectinload

from . import auth
from .ai import AiUnavailable, draft_reply
from .alerts import collect as collect_alerts, render_html as render_alert_html
from .config import ALERTS_KEY, EDITIONS, RECIPIENT_LISTS, settings
from .events import record
from .db import get_db, session_scope
from .models import (AiRule, ApiKey, Employee, Location, ReplyTemplate, ReportRecipient, Response as ReplyRow, Review,
                     ReviewEvent, ReviewMention, ReviewSourceLink, SavedView, SiteGroup, SyncRun, User)
from .daterange import PRESETS, DateRange, resolve_range
from .reports import (BRAND_COLORS, build_digest, build_trends, disputed_count, employee_report, listing_summaries, location_rank,
                      monthly_summary, rating_distribution, recovery_stats, render_digest_html, responder_stats, window_report)
from .sources import get_adapter
from .sync import sync_all
from .text_intel import THEMES as INTEL_THEMES, apply_intel, suggest_templates

log = logging.getLogger(__name__)
HERE = Path(__file__).parent

from contextlib import asynccontextmanager


@asynccontextmanager
async def _lifespan(_app: FastAPI):
    from .db import init_db
    init_db()          # create tables / add new columns before serving
    if settings.sso_config_problem:
        log.warning("microsoft sign-in disabled: %s", settings.sso_config_problem)
    yield


app = FastAPI(title="Review Manager", docs_url=None, redoc_url=None, lifespan=_lifespan)
app.mount("/static", StaticFiles(directory=str(HERE / "static")), name="static")
from .api import router as api_router, new_key as _new_api_key, hash_key as _hash_api_key  # noqa: E402
app.include_router(api_router)
if settings.api_cors_origins:
    from fastapi.middleware.cors import CORSMiddleware
    app.add_middleware(CORSMiddleware, allow_origins=settings.api_cors_origins, allow_methods=["GET"], allow_headers=["X-API-Key", "Authorization"])
templates = Jinja2Templates(directory=str(HERE / "templates"))

_tz = ZoneInfo(settings.timezone)
THEMES = INTEL_THEMES
TEMPLATE_TAGS = ["general", "no_comment", "team", "employee", "amenities", "membership", "mixed", "wait", "quality",
                 "damage", "billing", "staff", "plate", "hours", "vacuums"]
ATTENTION_WINDOW_DAYS = 90


def _local(dt: Optional[datetime], fmt: str = "%b %-d, %Y %-I:%M %p") -> str:
    if not dt:
        return ""
    return dt.replace(tzinfo=ZoneInfo("UTC")).astimezone(_tz).strftime(fmt)


def _age(dt: Optional[datetime]) -> str:
    if not dt:
        return ""
    delta = datetime.utcnow() - dt
    h = delta.total_seconds() / 3600
    if h < 1:
        return f"{max(1, int(delta.total_seconds() // 60))}m"
    if h < 48:
        return f"{int(h)}h"
    return f"{int(h // 24)}d"


def _hours(h: Optional[float]) -> str:
    if h is None:
        return "–"
    if h < 1:
        return f"{int(h * 60)}m"
    if h < 48:
        return f"{h:.0f}h"
    return f"{h / 24:.1f}d"


templates.env.filters["local"] = _local
templates.env.filters["age"] = _age
templates.env.filters["hours"] = _hours
templates.env.globals["settings"] = settings
templates.env.globals["brand_colors"] = BRAND_COLORS
templates.env.globals["now_utc"] = datetime.utcnow
templates.env.globals["THEMES"] = THEMES
templates.env.globals["PRESETS"] = PRESETS
templates.env.globals["EDITIONS"] = EDITIONS
templates.env.globals["RECIPIENT_LISTS"] = RECIPIENT_LISTS
templates.env.globals["ALERTS_KEY"] = ALERTS_KEY
templates.env.globals["GOOGLE_REPORT_TOOL"] = "https://support.google.com/business/workflow/9945796"


def _dr(range: str, start: str, end: str) -> DateRange:
    return resolve_range(range or "last30", start or "", end or "")


def _ints(values) -> List[int]:
    out: List[int] = []
    for v in values or []:
        for part in str(v).split(","):
            part = part.strip()
            if part and part != "0":
                try:
                    out.append(int(part))
                except ValueError:
                    pass
    return out


def _strs(values) -> List[str]:
    out: List[str] = []
    for v in values or []:
        for part in str(v).split(","):
            if part.strip():
                out.append(part.strip())
    return out


def _scope(db: Session, group_ids: List[int], location_ids: List[int]) -> Optional[List[int]]:
    """Resolve group + site multi-selects into one location id list (None = no restriction)."""
    ids: Optional[List[int]] = None
    if group_ids:
        ids = []
        for g in db.execute(select(SiteGroup).where(SiteGroup.id.in_(group_ids)).options(selectinload(SiteGroup.locations))).scalars().all():
            ids.extend(l.id for l in g.locations)
        ids = sorted(set(ids))
    if location_ids:
        ids = sorted(set(location_ids) if ids is None else set(ids) & set(location_ids))
    return ids


def _filter_ctx(db: Session) -> dict:
    return {"locations": db.execute(select(Location).where(Location.active.is_(True)).order_by(Location.brand, Location.name)).scalars().all(),
            "groups": db.execute(select(SiteGroup).order_by(SiteGroup.name)).scalars().all(), "brands": _brands(db)}


@app.middleware("http")
async def _renew_session(request: Request, call_next):
    """Sliding session: re-issue the cookie about once an hour while the person keeps using the app."""
    response = await call_next(request)
    uid = getattr(request.state, "renew_session_uid", None)
    if uid:
        response.set_cookie(auth.COOKIE_NAME, auth._serializer.dumps({"uid": uid}), httponly=True, samesite="lax",
                            max_age=auth.SESSION_MAX_AGE, secure=settings.app_base_url.startswith("https"))
    return response


@app.exception_handler(HTTPException)
async def _auth_redirect(request: Request, exc: HTTPException):
    if request.url.path.startswith("/api/") or "application/json" in (request.headers.get("accept") or ""):
        return JSONResponse({"error": exc.detail, "status": exc.status_code}, status_code=exc.status_code)
    if exc.status_code == status.HTTP_401_UNAUTHORIZED:
        return RedirectResponse(url=f"/login?next={request.url.path}", status_code=303)
    return HTMLResponse(f"<h1>{exc.status_code}</h1><p>{exc.detail}</p>", status_code=exc.status_code)


def _attention_cond(overdue_cut: datetime):
    """Negative and unanswered, or overdue within the last 90 days."""
    recent = datetime.utcnow() - timedelta(days=ATTENTION_WINDOW_DAYS)
    return or_(Review.rating <= settings.negative_rating_max,
               (Review.created_at_source < overdue_cut) & (Review.created_at_source >= recent))


def _open_base():
    return (select(func.count()).select_from(Review).join(ReviewSourceLink)
            .where(Review.is_deleted.is_(False), ReviewSourceLink.active.is_(True),
                   Review.has_owner_reply.is_(False), Review.is_archived.is_(False)))


def _sync_health(db: Session) -> dict:
    """Last sync, when the next one is due, and whether anything is failing. Used by the nav and /health."""
    last = db.execute(select(func.max(SyncRun.finished_at))).scalar()
    interval = timedelta(minutes=max(1, settings.sync_interval_minutes))
    now = datetime.utcnow()
    failing = db.execute(select(func.count()).select_from(ReviewSourceLink)
                         .where(ReviewSourceLink.active.is_(True), ReviewSourceLink.fail_count >= 1)).scalar() or 0
    next_due = (last + interval) if last else None
    stale = bool(last) and (now - last) > interval * 2
    return {"last_sync": last, "next_due": next_due, "next_in_min": max(0, int((next_due - now).total_seconds() // 60)) if next_due else None,
            "stale": stale, "failing": failing, "interval_min": settings.sync_interval_minutes}


def _nav_counts(db: Session) -> dict:
    overdue_cut = datetime.utcnow() - timedelta(hours=settings.overdue_hours)
    h = _sync_health(db)
    return {
        "attention": db.execute(_open_base().where(_attention_cond(overdue_cut))).scalar() or 0,
        "unanswered": db.execute(_open_base()).scalar() or 0,
        "last_sync": h["last_sync"], "sync": h,
    }


def _render(request: Request, name: str, user: Optional[User], db: Optional[Session] = None, **ctx) -> HTMLResponse:
    ctx.update({"request": request, "user": user, "nav": _nav_counts(db) if (db is not None and user) else {}})
    # Modern Starlette signature (request first); the legacy (name, context) form was removed in Starlette 0.50.
    return templates.TemplateResponse(request, name, ctx)


def _brands(db: Session) -> List[str]:
    return sorted({l.brand for l in db.execute(select(Location)).scalars().all()})


# ----------------------------------------------------------------- auth pages
@app.get("/login", response_class=HTMLResponse)
def login_form(request: Request, next: str = "/", error: str = ""):
    return _render(request, "login.html", None, next=next, error=error, sso_enabled=settings.sso_enabled,
                   password_enabled=settings.password_login_enabled, domains=settings.sso_allowed_domains)


def _session_redirect(user: User, next_url: str) -> RedirectResponse:
    resp = RedirectResponse(url=next_url or "/", status_code=303)
    resp.set_cookie(auth.COOKIE_NAME, auth.make_session_cookie(user), httponly=True, samesite="lax",
                    max_age=auth.SESSION_MAX_AGE, secure=settings.app_base_url.startswith("https"))
    return resp


@app.get("/auth/microsoft")
def microsoft_login(next: str = "/"):
    if not settings.sso_enabled:
        raise HTTPException(404, "Microsoft sign-in is not configured")
    auth_url, flow_cookie = auth.start_microsoft_flow()
    resp = RedirectResponse(url=auth_url, status_code=303)
    resp.set_cookie(auth.FLOW_COOKIE, flow_cookie, httponly=True, samesite="lax", max_age=auth.FLOW_MAX_AGE,
                    secure=settings.app_base_url.startswith("https"))
    resp.set_cookie("rm_next", next if next.startswith("/") else "/", httponly=True, samesite="lax", max_age=auth.FLOW_MAX_AGE)
    return resp


@app.get("/auth/microsoft/callback")
def microsoft_callback(request: Request, db: Session = Depends(get_db)):
    if not settings.sso_enabled:
        raise HTTPException(404, "Microsoft sign-in is not configured")
    try:
        claims = auth.finish_microsoft_flow(request.cookies.get(auth.FLOW_COOKIE), request.query_params)
        user = auth.resolve_sso_user(db, claims)
    except auth.SsoError as exc:
        log.warning("microsoft sign-in rejected: %s", exc)
        resp = RedirectResponse(url=f"/login?error={str(exc).replace(' ', '+')}", status_code=303)
        resp.delete_cookie(auth.FLOW_COOKIE)
        return resp
    next_url = request.cookies.get("rm_next") or "/"
    resp = _session_redirect(user, next_url if next_url.startswith("/") else "/")
    resp.delete_cookie(auth.FLOW_COOKIE)
    resp.delete_cookie("rm_next")
    return resp


@app.post("/login")
def login(request: Request, email: str = Form(...), password: str = Form(...), next: str = Form("/"),
          db: Session = Depends(get_db)):
    keys = (f"e:{email.strip().lower()}", f"ip:{request.client.host if request.client else '?'}")
    wait = auth.login_blocked(*keys)
    if wait:
        return _render(request, "login.html", None, next=next, error=f"Too many attempts. Try again in {max(1, wait // 60)} min.",
                       sso_enabled=settings.sso_enabled, password_enabled=settings.password_login_enabled, domains=settings.sso_allowed_domains)
    user = auth.authenticate(db, email, password)
    if not user:
        auth.note_login_failure(*keys)
        return _render(request, "login.html", None, next=next, error="Wrong email or password.", sso_enabled=settings.sso_enabled,
                       password_enabled=settings.password_login_enabled, domains=settings.sso_allowed_domains)
    auth.clear_login_failures(*keys)
    return _session_redirect(user, next if next.startswith("/") else "/")


@app.post("/logout")
def logout():
    resp = RedirectResponse(url="/login", status_code=303)
    resp.delete_cookie(auth.COOKIE_NAME)
    return resp


@app.get("/login/forgot", response_class=HTMLResponse)
def forgot_form(request: Request):
    return _render(request, "login_forgot.html", None, sent=False, email="", link="", error="")


@app.post("/login/forgot", response_class=HTMLResponse)
def forgot_submit(request: Request, email: str = Form(...), db: Session = Depends(get_db)):
    """Always answers the same way so addresses cannot be probed. Sends a 2-hour link when the
    account exists; shows the link on screen when email is not configured (dev / demo)."""
    from .account_mail import send_reset
    email = email.strip().lower()
    key = f"reset:{email}"
    if auth.login_blocked(key):
        return _render(request, "login_forgot.html", None, sent=False, email=email, link="", error="Too many requests. Try again in a few minutes.")
    auth.note_login_failure(key)
    link = ""
    if settings.password_login_enabled:
        u = db.execute(select(User).where(User.email == email)).scalar_one_or_none()
        if u and u.active:
            link = auth.password_link(u, "reset")
            if send_reset(u, link):
                link = ""
            elif settings.mail_enabled:
                link = ""      # sending failed; never show a live link when mail was supposed to work
    return _render(request, "login_forgot.html", None, sent=True, email=email, link=link, error="")


@app.get("/password/reset", response_class=HTMLResponse)
def reset_form(request: Request, token: str = "", db: Session = Depends(get_db)):
    target = auth.verify_password_token(db, token)
    purpose = "welcome" if (target and not target.password_hash) else "reset"
    return _render(request, "set_password.html", None, target=target, token=token, purpose=purpose, error="")


@app.post("/password/reset")
def reset_submit(request: Request, token: str = Form(...), password: str = Form(...), confirm: str = Form(""), db: Session = Depends(get_db)):
    target = auth.verify_password_token(db, token)
    if not target:
        return _render(request, "set_password.html", None, target=None, token=token, purpose="reset", error="")
    problem = auth.password_problem(password, confirm)
    if problem:
        return _render(request, "set_password.html", None, target=target, token=token, purpose="welcome" if not target.password_hash else "reset", error=problem)
    target.password_hash = auth.hash_password(password)
    target.last_login_at = datetime.utcnow()
    db.commit()
    auth.clear_login_failures(f"e:{target.email}", f"reset:{target.email}")
    return _session_redirect(target, "/?msg=Password+saved.+You+are+signed+in.")


@app.get("/account", response_class=HTMLResponse)
def account(request: Request, user: User = Depends(auth.current_user), db: Session = Depends(get_db), error: str = ""):
    return _render(request, "account.html", user, db, error=error, password_enabled=settings.password_login_enabled)


@app.post("/account/password")
def change_password(user: User = Depends(auth.current_user), db: Session = Depends(get_db), current: str = Form(""),
                    password: str = Form(...), confirm: str = Form("")):
    if not settings.password_login_enabled:
        return RedirectResponse(url="/account?error=Password+sign-in+is+turned+off", status_code=303)
    if user.password_hash and not auth.verify_password(current, user.password_hash):
        return RedirectResponse(url=f"/account?{urlencode({'error': 'The current password is wrong.'})}", status_code=303)
    problem = auth.password_problem(password, confirm)
    if problem:
        return RedirectResponse(url=f"/account?{urlencode({'error': problem})}", status_code=303)
    user.password_hash = auth.hash_password(password)
    db.commit()
    return RedirectResponse(url="/account?msg=Password+saved", status_code=303)


@app.get("/health")
def health(db: Session = Depends(get_db)):
    h = _sync_health(db)
    return {"ok": not h["stale"] and not h["failing"], "last_sync_finished_at": h["last_sync"].isoformat() if h["last_sync"] else None,
            "next_sync_due": h["next_due"].isoformat() if h["next_due"] else None, "stale": h["stale"], "failing_listings": h["failing"]}


# ----------------------------------------------------------------- inbox
VIEWS = [("attention", "Needs attention"), ("unanswered", "Unanswered"), ("negative", "Negative"),
         ("replied", "Replied"), ("all", "All"), ("failed", "Failed"), ("archived", "Archived"), ("removed", "Removed")]
VIEW_KEYS = {k for k, _ in VIEWS}


def _inbox_query(user: User, view: str, brands: List[str], scope_ids: Optional[List[int]], ratings: List[int], q: str, dr: Optional[DateRange],
                 overdue_cut: datetime):
    base = (select(Review)
            .join(ReviewSourceLink, Review.source_link_id == ReviewSourceLink.id)
            .outerjoin(Location, ReviewSourceLink.location_id == Location.id)
            .where(Review.is_deleted.is_(view == "removed"), ReviewSourceLink.active.is_(True))
            .options(selectinload(Review.source_link).selectinload(ReviewSourceLink.location),
                     selectinload(Review.assigned_to), selectinload(Review.mentions),
                     selectinload(Review.responses).selectinload(ReplyRow.created_by)))
    order = (Review.removed_at.desc(),) if view == "removed" else (Review.created_at_source.desc(),)
    open_ = (Review.has_owner_reply.is_(False)) & (Review.is_archived.is_(False))
    if view == "attention":
        base = base.where(open_, _attention_cond(overdue_cut))
        order = ((Review.rating <= settings.negative_rating_max).desc(), Review.created_at_source.asc())
    elif view == "unanswered":
        base = base.where(open_)
        order = (Review.created_at_source.asc(),)
    elif view == "negative":
        base = base.where(Review.rating <= settings.negative_rating_max)
    elif view == "replied":
        base = base.where(Review.has_owner_reply.is_(True))
    elif view == "failed":
        failed_ids = select(ReplyRow.review_id).where(ReplyRow.status == "failed")
        base = base.where(open_, Review.id.in_(failed_ids))
    elif view == "archived":
        base = base.where(Review.is_archived.is_(True))
    if brands:
        base = base.where(Location.brand.in_(brands))
    if scope_ids is not None:
        base = base.where(Location.id.in_(scope_ids or [-1]))
    if ratings:
        base = base.where(Review.rating.in_(ratings))
    if dr is not None:
        base = base.where(Review.created_at_source >= dr.start, Review.created_at_source < dr.end)
    if q:
        like = f"%{q}%"
        base = base.where(or_(Review.text.ilike(like), Review.author_name.ilike(like)))
    return base, order


@app.get("/", response_class=HTMLResponse)
def inbox(request: Request, user: User = Depends(auth.current_user), db: Session = Depends(get_db),
          view: str = "attention", brand: List[str] = Query([]), location_id: List[str] = Query([]), group_id: List[str] = Query([]),
          rating: List[str] = Query([]), q: str = "", range: str = "", start: str = "", end: str = "", sv: int = 0, page: int = 1,
          sort: str = "", dir: str = "asc"):
    per_page = 50
    brands_f, loc_f, grp_f, rat_f = _strs(brand), _ints(location_id), _ints(group_id), _ints(rating)
    saved_views = db.execute(select(SavedView).where(or_(SavedView.is_shared.is_(True), SavedView.owner_id == user.id))
                             .order_by(SavedView.sort_order, SavedView.name)).scalars().all()
    active_view = None
    if sv:
        active_view = next((v for v in saved_views if v.id == sv), None)
        if active_view:
            p = active_view.params()
            view = p.get("view", view); q = p.get("q", "")
            brands_f = _strs(p.get("brands") or ([p["brand"]] if p.get("brand") else []))
            loc_f = _ints(p.get("location_ids") or ([p["location_id"]] if p.get("location_id") else []))
            grp_f = _ints(p.get("group_ids") or ([p["group_id"]] if p.get("group_id") else []))
            rat_f = _ints(p.get("ratings") or ([p["rating"]] if p.get("rating") else []))
            range = p.get("range", ""); start = p.get("start", ""); end = p.get("end", "")
            if not range and p.get("days"):
                range = {1: "yesterday", 7: "last7", 30: "last30"}.get(int(p["days"]), "")
    if view not in VIEW_KEYS:
        view = "attention"
    overdue_cut = datetime.utcnow() - timedelta(hours=settings.overdue_hours)
    scope_ids = _scope(db, grp_f, loc_f)
    dr = _dr(range, start, end) if (range or start or end) else None
    base, order = _inbox_query(user, view, brands_f, scope_ids, rat_f, q, dr, overdue_cut)
    dir = "desc" if dir == "desc" else "asc"
    sort_cols = {
        "rating": (Review.rating,), "site": (Location.name,), "author": (Review.author_name,), "created": (Review.created_at_source,),
        "status": (Review.has_owner_reply, Review.rating, Review.created_at_source),
    }
    if sort in sort_cols:
        order = tuple((c.desc() if dir == "desc" else c.asc()) for c in sort_cols[sort]) + (Review.created_at_source.desc(),)
    else:
        sort = ""
    total = db.execute(select(func.count()).select_from(base.order_by(None).subquery())).scalar() or 0
    rows = db.execute(base.order_by(*order).offset((page - 1) * per_page).limit(per_page)).scalars().all()

    cbase = _open_base()
    all_base = (select(func.count()).select_from(Review).join(ReviewSourceLink)
                .where(Review.is_deleted.is_(False), ReviewSourceLink.active.is_(True)))
    counts = {
        "attention": db.execute(cbase.where(_attention_cond(overdue_cut))).scalar(),
        "unanswered": db.execute(cbase).scalar(),
        "negative": db.execute(all_base.where(Review.rating <= settings.negative_rating_max)).scalar(),
        "replied": db.execute(all_base.where(Review.has_owner_reply.is_(True))).scalar(),
        "failed": db.execute(cbase.where(Review.id.in_(select(ReplyRow.review_id).where(ReplyRow.status == "failed")))).scalar(),
        "removed": db.execute(select(func.count()).select_from(Review).join(ReviewSourceLink)
                              .where(Review.is_deleted.is_(True), ReviewSourceLink.active.is_(True))).scalar(),
    }
    current_params = {"view": view, "brands": brands_f, "location_ids": loc_f, "group_ids": grp_f, "ratings": rat_f, "q": q,
                      "range": dr.preset if dr else "", "start": dr.start_date.isoformat() if (dr and dr.is_custom) else "",
                      "end": dr.end_date.isoformat() if (dr and dr.is_custom) else ""}
    qs_parts = [f"view={view}", f"q={q}", f"range={current_params['range']}", f"start={current_params['start']}", f"end={current_params['end']}"]
    qs_parts += [f"brand={b}" for b in brands_f] + [f"location_id={i}" for i in loc_f] + [f"group_id={i}" for i in grp_f] + [f"rating={i}" for i in rat_f]
    qs = "&".join(qs_parts)
    filtered = bool(q or brands_f or loc_f or grp_f or rat_f or dr)
    # Queue context: the review page uses it for prev/next within exactly this list.
    ctx = f"{qs}&sort={sort}&dir={dir}&sv={active_view.id if active_view else 0}"
    # Inline composer: top template suggestions per open review on this page.
    tpl_rows = db.execute(select(ReplyTemplate).where(ReplyTemplate.active.is_(True)).order_by(ReplyTemplate.sort_order, ReplyTemplate.name)).scalars().all()
    first = user.name.split()[0] if user.name else ""
    inline = {}
    for r in rows:
        if r.has_owner_reply or r.is_archived or r.is_deleted:
            continue
        ranked = suggest_templates(r, tpl_rows, r.mention_names)[:4]
        inline[r.id] = [{"id": t.id, "name": t.name, "body": t.render(r, first), "suggested": i < 3 and sc > 0} for i, (t, sc) in enumerate(ranked)]
    changed_since = datetime.utcnow() - timedelta(days=30)
    return _render(request, "inbox.html", user, db, reviews=rows, total=total, page=page, per_page=per_page,
                   view=view, views=VIEWS, brands_f=brands_f, loc_f=loc_f, grp_f=grp_f, rat_f=rat_f, q=q, dr=dr, filtered=filtered,
                   counts=counts, overdue_cut=overdue_cut, saved_views=saved_views, active_view=active_view, qs=qs, ctx=ctx,
                   inline_json=json.dumps(inline), inline=inline, changed_since=changed_since, ai_enabled=settings.ai_enabled,
                   current_params_json=json.dumps(current_params), sort=sort, dir=dir, **_filter_ctx(db))


@app.post("/groups")
def create_group(request: Request, user: User = Depends(auth.current_user), db: Session = Depends(get_db), name: str = Form(...),
                 location_ids: List[int] = Form([]), back: str = Form("/")):
    name = name.strip()[:80]
    if not name or not location_ids:
        return RedirectResponse(url=back or "/", status_code=303)
    g = db.execute(select(SiteGroup).where(SiteGroup.name == name)).scalar_one_or_none() or SiteGroup(name=name)
    db.add(g)
    g.locations = db.execute(select(Location).where(Location.id.in_(location_ids))).scalars().all()
    db.commit()
    sep = "&" if "?" in back else "?"
    return RedirectResponse(url=f"{back}{sep}group_id={g.id}" if back.startswith("/") else f"/?group_id={g.id}", status_code=303)


@app.post("/views")
def save_view(user: User = Depends(auth.current_user), db: Session = Depends(get_db), name: str = Form(...),
              is_shared: int = Form(0), params: str = Form("{}")):
    try:
        p = json.loads(params)
    except ValueError:
        p = {}
    v = SavedView(name=name.strip()[:80] or "Untitled", owner_id=user.id, is_shared=bool(is_shared), params_json=json.dumps(p))
    db.add(v)
    db.commit()
    return RedirectResponse(url=f"/?sv={v.id}", status_code=303)


@app.post("/views/{vid}/delete")
def delete_view(vid: int, user: User = Depends(auth.current_user), db: Session = Depends(get_db)):
    v = db.get(SavedView, vid)
    if v and (v.owner_id == user.id or user.is_admin):
        db.delete(v)
        db.commit()
    return RedirectResponse(url="/", status_code=303)


# ----------------------------------------------------------------- review detail / reply
def _get_review(db: Session, review_id: int) -> Review:
    r = db.execute(select(Review).where(Review.id == review_id)
                   .options(selectinload(Review.source_link).selectinload(ReviewSourceLink.location),
                            selectinload(Review.responses).selectinload(ReplyRow.created_by),
                            selectinload(Review.assigned_to), selectinload(Review.mentions).selectinload(ReviewMention.employee))).scalar_one_or_none()
    if not r:
        raise HTTPException(404, "Review not found")
    return r


def _queue_ids(db: Session, user: User, ctx: str) -> List[int]:
    """Review ids in the exact order of the inbox list the person came from."""
    p = parse_qs(ctx, keep_blank_values=True)
    g = lambda k, d="": (p.get(k) or [d])[0]
    view = g("view", "attention")
    if view not in VIEW_KEYS:
        view = "attention"
    dr = _dr(g("range"), g("start"), g("end")) if (g("range") or g("start") or g("end")) else None
    overdue_cut = datetime.utcnow() - timedelta(hours=settings.overdue_hours)
    base, order = _inbox_query(user, view, _strs(p.get("brand")), _scope(db, _ints(p.get("group_id")), _ints(p.get("location_id"))),
                               _ints(p.get("rating")), g("q"), dr, overdue_cut)
    sort, direction = g("sort"), g("dir", "asc")
    sort_cols = {"rating": (Review.rating,), "site": (Location.name,), "author": (Review.author_name,), "created": (Review.created_at_source,)}
    if sort in sort_cols:
        order = tuple((c.desc() if direction == "desc" else c.asc()) for c in sort_cols[sort]) + (Review.created_at_source.desc(),)
    return list(db.execute(base.with_only_columns(Review.id).order_by(None).order_by(*order).limit(2000)).scalars().all())


def _neighbors(db: Session, r: Review, user: Optional[User] = None, ctx: str = "") -> dict:
    """Prev/next within the queue the person is working (the inbox list they came from),
    falling back to the open-reviews queue in time order."""
    if ctx and user is not None:
        ids = _queue_ids(db, user, ctx)
        if r.id in ids:
            i = ids.index(r.id)
            return {"prev": ids[i - 1] if i > 0 else None, "next": ids[i + 1] if i + 1 < len(ids) else None,
                    "pos": i + 1, "total": len(ids), "ctx": ctx}
        return {"prev": None, "next": ids[0] if ids else None, "pos": None, "total": len(ids), "ctx": ctx}
    open_ = (Review.is_deleted.is_(False)) & (ReviewSourceLink.active.is_(True)) & (Review.has_owner_reply.is_(False)) & (Review.is_archived.is_(False))
    nxt = db.execute(select(Review.id).join(ReviewSourceLink).where(open_, Review.created_at_source > r.created_at_source)
                     .order_by(Review.created_at_source.asc()).limit(1)).scalar()
    prv = db.execute(select(Review.id).join(ReviewSourceLink).where(open_, Review.created_at_source < r.created_at_source)
                     .order_by(Review.created_at_source.desc()).limit(1)).scalar()
    return {"next": nxt, "prev": prv, "pos": None, "total": None, "ctx": ""}


PLACEHOLDER_RE = re.compile(r"\{[A-Za-z_ ]+\}|\$\{[^}]+\}|\[[^\]\n]{1,40}\]")
GREETING_RE = re.compile(r"^\s*(?:hi|hello|hey|dear|thanks|thank you)[,!]?\s+([A-Z][a-z]+)\b", re.IGNORECASE)


def _reply_warnings(db: Session, r: Review, text: str) -> List[str]:
    """Things worth a second look before a reply goes public. Never blocks; the UI offers 'post anyway'."""
    w: List[str] = []
    found = sorted(set(PLACEHOLDER_RE.findall(text)))
    if found:
        w.append("Unfilled placeholder still in the reply: " + ", ".join(found[:3]))
    if len(text) > 4096:
        w.append(f"Too long for Google: {len(text)} of 4096 characters")
    m = GREETING_RE.match(text)
    if m and r.first_name != "there" and m.group(1).lower() != r.first_name.lower():
        w.append(f"The reply greets {m.group(1)} but this reviewer is {r.first_name}")
    norm = " ".join(text.lower().split())
    since = datetime.utcnow() - timedelta(hours=24)
    others = db.execute(select(ReplyRow.text).join(Review, ReplyRow.review_id == Review.id)
                        .where(ReplyRow.status == "posted", ReplyRow.posted_at >= since,
                               Review.source_link_id == r.source_link_id, Review.id != r.id)).scalars().all()
    dup = sum(1 for t in others if " ".join((t or "").lower().split()) == norm)
    if dup >= 2:
        w.append(f"This exact reply was already posted {dup} times today at this site. Google may flag repeated text.")
    return w


def _timeline(r: Review) -> List[Dict]:
    """Events, newest first. Older data has no events, so reply rows stand in for them."""
    items = [{"at": e.at, "label": e.label, "who": e.actor.name if e.actor else "Sync", "kind": e.kind, "detail": e.detail} for e in r.events]
    if not any(i["kind"].startswith("reply") for i in items):
        for h in r.responses:
            label = {"posted": "Reply posted", "failed": "Reply failed to post", "deleted": "Reply removed"}.get(h.status, h.status)
            items.append({"at": h.created_at, "label": label, "who": h.created_by.name if h.created_by else "—", "kind": "reply_" + h.status,
                          "detail": {"text": h.text, "error": h.error, "ai": h.ai_generated, "template": bool(h.template_id)}})
    items.sort(key=lambda x: x["at"], reverse=True)
    return items


def _template_suggestions(db: Session, r: Review, user: User) -> List[Dict]:
    tpls = db.execute(select(ReplyTemplate).where(ReplyTemplate.active.is_(True)).order_by(ReplyTemplate.sort_order, ReplyTemplate.name)).scalars().all()
    ranked = suggest_templates(r, tpls, r.mention_names)
    first = user.name.split()[0] if user.name else ""
    out = []
    for i, (t, score) in enumerate(ranked):
        out.append({"id": t.id, "name": t.name, "body": t.render(r, first), "score": score, "suggested": i < 3 and score > 0,
                    "band": t.rating_band, "usage": t.usage_count or 0})
    return out


def _review_page(request: Request, user: User, db: Session, r: Review, ctx: str = "", msg: str = "",
                 warnings: Optional[List[str]] = None, draft: Optional[str] = None, status_code: int = 200) -> HTMLResponse:
    tpls = _template_suggestions(db, r, user)
    others = db.execute(select(Review).join(ReviewSourceLink).where(
        Review.source_link_id == r.source_link_id, Review.id != r.id, Review.is_deleted.is_(False),
        Review.author_name == r.author_name, Review.author_name.isnot(None)).order_by(Review.created_at_source.desc()).limit(5)).scalars().all() if r.author_name else []
    resp = _render(request, "review.html", user, db, r=r, msg=msg, templates_json=json.dumps(tpls), tpls=tpls, ctx=ctx,
                   nav_links=_neighbors(db, r, user, ctx), same_author=others, ai_enabled=settings.ai_enabled,
                   warnings=warnings or [], draft=draft, timeline=_timeline(r), listing=listing_summaries(db).get(r.location.id) if r.location else None)
    resp.status_code = status_code
    return resp


@app.get("/reviews/{review_id}", response_class=HTMLResponse)
def review_detail(review_id: int, request: Request, user: User = Depends(auth.current_user),
                  db: Session = Depends(get_db), msg: str = "", ctx: str = ""):
    return _review_page(request, user, db, _get_review(db, review_id), ctx=ctx, msg=msg)


def _wants_json(request: Request) -> bool:
    return "application/json" in (request.headers.get("accept") or "")


def _reply_cell(r: Review, by: str) -> Dict:
    return {"text": r.owner_reply_text or "", "by": by, "when": _local(r.owner_reply_updated_at, "%b %-d"),
            "hours": _hours(r.response_hours) if r.response_hours is not None else ""}


@app.post("/reviews/{review_id}/reply/check")
def reply_check(review_id: int, text: str = Form(""), user: User = Depends(auth.current_user), db: Session = Depends(get_db)):
    """Pre-flight guardrails for the composer; returns warnings without posting anything."""
    r = _get_review(db, review_id)
    return JSONResponse({"warnings": _reply_warnings(db, r, text.strip())})


@app.post("/reviews/{review_id}/reply")
def post_reply(review_id: int, request: Request, text: str = Form(...), go_next: int = Form(0), template_id: int = Form(0),
               ai_generated: int = Form(0), ctx: str = Form(""), force: int = Form(0),
               user: User = Depends(auth.current_user), db: Session = Depends(get_db)):
    """Create or replace our reply. Works for the page form and for the inbox's inline composer
    (Accept: application/json). Guardrail warnings come back first unless force=1."""
    r = _get_review(db, review_id)
    text = text.strip()
    wants_json = _wants_json(request)
    if r.is_deleted:
        msg = "This review is no longer on the platform, so a reply cannot be posted."
        return JSONResponse({"ok": False, "error": msg}, status_code=409) if wants_json else _review_page(request, user, db, r, ctx, msg=msg)
    if not text:
        return JSONResponse({"ok": False, "error": "Reply was empty"}, status_code=400) if wants_json else \
            RedirectResponse(url=f"/reviews/{review_id}?msg=Reply+was+empty&ctx={ctx}", status_code=303)
    warnings = _reply_warnings(db, r, text) if not force else []
    if warnings:
        if wants_json:
            return JSONResponse({"ok": False, "warnings": warnings}, status_code=422)
        return _review_page(request, user, db, r, ctx, warnings=warnings, draft=text)
    nxt = _neighbors(db, r, user, ctx)["next"] if go_next else None    # before the state changes
    editing = r.has_owner_reply
    row = ReplyRow(review_id=r.id, text=text, created_by_id=user.id, status="draft",
                   template_id=template_id or None, ai_generated=bool(ai_generated))
    db.add(row)
    db.flush()
    try:
        posted_at = get_adapter(r.source).post_reply(r.source_link, r.external_id, text)
        row.status, row.posted_at = "posted", posted_at
        r.has_owner_reply, r.owner_reply_text, r.owner_reply_updated_at = True, text, posted_at
        if r.first_replied_at is None:
            r.first_replied_at = posted_at
        r.assigned_to_id = r.assigned_to_id or user.id
        if template_id:
            t = db.get(ReplyTemplate, template_id)
            if t:
                t.usage_count = (t.usage_count or 0) + 1
                t.last_used_at = datetime.utcnow()
        record(db, r, "reply_edited" if editing else "reply_posted", actor_id=user.id, at=posted_at, text=text,
               template=bool(template_id), ai=bool(ai_generated))
        msg = "Reply updated." if editing else "Reply posted."
        ok = True
    except Exception as exc:
        log.exception("reply failed")
        row.status, row.error = "failed", str(exc)[:2000]
        record(db, r, "reply_failed", actor_id=user.id, error=str(exc)[:300])
        msg, ok = "Posting failed. See the error below.", False
    db.commit()
    if wants_json:
        return JSONResponse({"ok": ok, "msg": msg, "reply": _reply_cell(r, user.name.split()[0]) if ok else None,
                             "next": nxt if ok else None, "error": row.error if not ok else None}, status_code=200 if ok else 502)
    if ok and nxt:
        return RedirectResponse(url=f"/reviews/{nxt}?msg={msg.replace(' ', '+')}+Here+is+the+next+one.&ctx={ctx}", status_code=303)
    return RedirectResponse(url=f"/reviews/{review_id}?msg={msg.replace(' ', '+')}&ctx={ctx}", status_code=303)


@app.post("/reviews/{review_id}/reply/delete")
def delete_reply(review_id: int, user: User = Depends(auth.current_user), db: Session = Depends(get_db)):
    r = _get_review(db, review_id)
    try:
        get_adapter(r.source).delete_reply(r.source_link, r.external_id)
        r.has_owner_reply = False
        r.owner_reply_text = None
        r.owner_reply_updated_at = None
        db.add(ReplyRow(review_id=r.id, text="(reply removed)", created_by_id=user.id, status="deleted", posted_at=datetime.utcnow()))
        record(db, r, "reply_deleted", actor_id=user.id)
        msg = "Reply removed."
    except Exception as exc:
        log.exception("delete reply failed")
        db.add(ReplyRow(review_id=r.id, text="(delete attempt)", created_by_id=user.id, status="failed", error=str(exc)[:2000]))
        record(db, r, "reply_failed", actor_id=user.id, error=str(exc)[:300], action="delete")
        msg = "Delete failed."
    db.commit()
    return RedirectResponse(url=f"/reviews/{review_id}?msg={msg.replace(' ', '+')}", status_code=303)


@app.post("/reviews/{review_id}/note")
def save_note(review_id: int, note: str = Form(""), category: str = Form(""), user: User = Depends(auth.current_user), db: Session = Depends(get_db)):
    r = _get_review(db, review_id)
    new_note, new_cat = note.strip() or None, category.strip() or None
    if new_cat != r.category:
        record(db, r, "category_changed", actor_id=user.id, **{"from": r.category, "to": new_cat})
        r.category, r.category_source = new_cat, ("manual" if new_cat else None)
    if new_note != r.internal_note:
        record(db, r, "note_saved", actor_id=user.id, text=(new_note or "")[:300])
        r.internal_note = new_note
    db.commit()
    return RedirectResponse(url=f"/reviews/{review_id}?msg=Saved", status_code=303)


@app.post("/reviews/{review_id}/claim")
def claim(review_id: int, user: User = Depends(auth.current_user), db: Session = Depends(get_db), back: str = Form("")):
    r = _get_review(db, review_id)
    if r.assigned_to_id == user.id:
        r.assigned_to_id, r.assigned_at = None, None
    else:
        r.assigned_to_id, r.assigned_at = user.id, datetime.utcnow()
    db.commit()
    return RedirectResponse(url=back or f"/reviews/{review_id}", status_code=303)


@app.post("/reviews/{review_id}/archive")
def archive(review_id: int, request: Request, user: User = Depends(auth.current_user), db: Session = Depends(get_db), back: str = Form("")):
    """Toggle archived. The redirect carries an Undo link; the inbox shortcut uses the JSON form."""
    r = _get_review(db, review_id)
    r.is_archived = not r.is_archived
    r.archived_at = datetime.utcnow() if r.is_archived else None
    r.archived_by_id = user.id if r.is_archived else None
    record(db, r, "archived" if r.is_archived else "unarchived", actor_id=user.id)
    db.commit()
    msg = "Archived. It no longer counts as unanswered." if r.is_archived else "Restored to the inbox."
    if _wants_json(request):
        return JSONResponse({"ok": True, "archived": r.is_archived, "msg": msg, "undo": f"/reviews/{r.id}/archive"})
    q = urlencode({"msg": msg, "undo": f"/reviews/{r.id}/archive"})
    if back and back.startswith("/"):
        return RedirectResponse(url=f"{back}{'&' if '?' in back else '?'}{q}", status_code=303)
    return RedirectResponse(url=f"/reviews/{review_id}?{q}", status_code=303)


@app.post("/reviews/{review_id}/report")
def report_review(review_id: int, user: User = Depends(auth.current_user), db: Session = Depends(get_db),
                  action: str = Form(...), note: str = Form("")):
    """Track a removal report made to the platform: reported -> removed / kept, or clear it."""
    r = _get_review(db, review_id)
    now = datetime.utcnow()
    if action == "reported":
        r.report_status, r.reported_at, r.reported_by_id = "reported", now, user.id
        r.report_note = note.strip() or r.report_note
        record(db, r, "reported", actor_id=user.id, note=note.strip()[:300])
        msg = "Marked as reported to Google. It is left out of averages until Google decides."
    elif action in ("removed", "kept"):
        r.report_status = action
        r.report_note = note.strip() or r.report_note
        record(db, r, "report_outcome", actor_id=user.id, outcome=action, note=note.strip()[:300])
        msg = "Outcome recorded: Google removed it." if action == "removed" else "Outcome recorded: Google kept it. It counts in averages again."
    else:
        r.report_status, r.reported_at, r.reported_by_id, r.report_note = None, None, None, None
        record(db, r, "report_outcome", actor_id=user.id, outcome="cleared")
        msg = "Report cleared."
    db.commit()
    return RedirectResponse(url=f"/reviews/{review_id}?{urlencode({'msg': msg})}", status_code=303)


@app.post("/reviews/{review_id}/mentions")
def add_mention(review_id: int, name: str = Form(...), user: User = Depends(auth.current_user), db: Session = Depends(get_db)):
    r = _get_review(db, review_id)
    name = name.strip()[:80]
    if name and name not in r.mention_names:
        emp = db.execute(select(Employee).where(func.lower(Employee.name) == name.lower())).scalar_one_or_none()
        db.add(ReviewMention(review_id=r.id, name=emp.name if emp else name, employee_id=emp.id if emp else None, source="manual"))
        record(db, r, "mention_added", actor_id=user.id, name=emp.name if emp else name)
        db.commit()
    return RedirectResponse(url=f"/reviews/{review_id}?msg=Mention+added", status_code=303)


@app.post("/reviews/{review_id}/mentions/{mid}/delete")
def delete_mention(review_id: int, mid: int, user: User = Depends(auth.current_user), db: Session = Depends(get_db)):
    m = db.get(ReviewMention, mid)
    if m and m.review_id == review_id:
        record(db, m.review, "mention_removed", actor_id=user.id, name=m.name)
        db.delete(m)
        db.commit()
    return RedirectResponse(url=f"/reviews/{review_id}?msg=Mention+removed", status_code=303)


@app.post("/reviews/{review_id}/draft")
def draft(review_id: int, user: User = Depends(auth.current_user), db: Session = Depends(get_db)):
    r = _get_review(db, review_id)
    rules = db.execute(select(AiRule).where(AiRule.active.is_(True)).order_by(AiRule.sort_order)).scalars().all()
    examples = [t for t, _ in suggest_templates(r, db.execute(select(ReplyTemplate).where(ReplyTemplate.active.is_(True))).scalars().all(), r.mention_names)][:5]
    try:
        text = draft_reply(r, rules, examples, user.name.split()[0] if user.name else "")
        return JSONResponse({"ok": True, "text": text})
    except AiUnavailable as exc:
        return JSONResponse({"ok": False, "error": str(exc)}, status_code=503)
    except Exception as exc:  # pragma: no cover
        log.exception("draft failed")
        return JSONResponse({"ok": False, "error": "Drafting failed"}, status_code=500)


# ----------------------------------------------------------------- reports
@app.get("/reports", response_class=HTMLResponse)
def reports(request: Request, user: User = Depends(auth.current_user), db: Session = Depends(get_db), brand: List[str] = Query([]),
            group_id: List[str] = Query([]), location_id: List[str] = Query([]), range: str = "last30", start: str = "", end: str = "",
            unit: str = ""):
    dr = _dr(range, start, end)
    brands_f, grp_f, loc_f = _strs(brand), _ints(group_id), _ints(location_id)
    ids = _scope(db, grp_f, loc_f)
    bl = brands_f or None
    wr = window_report(db, dr, brands=bl, location_ids=ids)
    trends = build_trends(db, dr, brands=bl, location_ids=ids, unit=unit or None)
    fq = "&".join([dr.query()] + [f"brand={b}" for b in brands_f] + [f"group_id={i}" for i in grp_f] + [f"location_id={i}" for i in loc_f])
    rec = recovery_stats(db, dr, brands=bl, location_ids=ids)
    responders = responder_stats(db, dr)
    for p_ in responders:
        p_["recovered"] = rec["by_responder"].get(p_["name"], 0)
    listing = listing_summaries(db)
    rank = location_rank(db, dr, brands=bl, location_ids=ids)
    for row in rank:
        ls = listing.get(row["location_id"]) if row["location_id"] else None
        row["listing_avg"], row["listing_total"] = (ls["avg"], ls["total"]) if ls else (None, None)
    return _render(request, "reports.html", user, db, w=wr, dr=dr, trends=trends, trends_json=json.dumps(trends), unit=trends["unit"],
                   brands_f=brands_f, grp_f=grp_f, loc_f=loc_f, fq=fq, recovery=rec, disputed=disputed_count(db, dr, bl, ids),
                   monthly=monthly_summary(db, bl, months=12, location_ids=ids), responders=responders,
                   rank=rank, ai_enabled=settings.ai_enabled, **_filter_ctx(db))


@app.get("/reports/distribution", response_class=HTMLResponse)
def distribution_report(request: Request, user: User = Depends(auth.current_user), db: Session = Depends(get_db), brand: List[str] = Query([]),
                        group_id: List[str] = Query([]), location_id: List[str] = Query([]), range: str = "last30", start: str = "", end: str = ""):
    dr = _dr(range, start, end)
    brands_f, grp_f, loc_f = _strs(brand), _ints(group_id), _ints(location_id)
    ids = _scope(db, grp_f, loc_f)
    dist = rating_distribution(db, dr, brands=brands_f or None, location_ids=ids)
    fq = "&".join([dr.query()] + [f"brand={b}" for b in brands_f] + [f"group_id={i}" for i in grp_f] + [f"location_id={i}" for i in loc_f])
    return _render(request, "reports_distribution.html", user, db, dist=dist, dist_json=json.dumps(dist["chart"]), dr=dr,
                   brands_f=brands_f, grp_f=grp_f, loc_f=loc_f, fq=fq, disputed=disputed_count(db, dr, brands_f or None, ids), **_filter_ctx(db))


@app.get("/reports/weekly", response_class=HTMLResponse)
def weekly_report(request: Request, user: User = Depends(auth.current_user), db: Session = Depends(get_db), brand: List[str] = Query([]),
                  group_id: List[str] = Query([]), location_id: List[str] = Query([]), as_of: str = ""):
    """Last week (Mon–Sun) and quarter-to-date side by side: reviews by rating per site, then
    negative review reasons per site. `as_of` moves the anchor date (default today)."""
    brands_f, grp_f, loc_f = _strs(brand), _ints(group_id), _ints(location_id)
    ids = _scope(db, grp_f, loc_f)
    bl = brands_f or None
    anchor = None
    try:
        anchor = datetime.fromisoformat(as_of).replace(tzinfo=_tz) if as_of else None
    except ValueError:
        anchor = None
    lw, qtd = resolve_range("last_week", now_local=anchor), resolve_range("this_quarter", now_local=anchor)
    dist_lw, dist_q = rating_distribution(db, lw, bl, ids), rating_distribution(db, qtd, bl, ids)
    w_lw, w_q = window_report(db, lw, bl, ids), window_report(db, qtd, bl, ids)
    names = sorted({r["name"] for r in dist_lw["rows"]} | {r["name"] for r in dist_q["rows"]} | set(w_lw.theme_matrix) | set(w_q.theme_matrix))
    by_lw, by_q = {r["name"]: r for r in dist_lw["rows"]}, {r["name"]: r for r in dist_q["rows"]}
    brand_of = {r["name"]: r["brand"] for r in dist_lw["rows"] + dist_q["rows"]}
    sites = sorted(names, key=lambda n: (["WashU", "ICON", "WA"].index(brand_of.get(n, "?")) if brand_of.get(n) in ["WashU", "ICON", "WA"] else 9, n))
    themes_used = [c for c in THEMES if any(w.theme_matrix.get(sname, {}).get(c) for w in (w_lw, w_q) for sname in sites)]
    fq = "&".join([f"brand={b}" for b in brands_f] + [f"group_id={i}" for i in grp_f] + [f"location_id={i}" for i in loc_f] + ([f"as_of={as_of}"] if as_of else []))
    return _render(request, "reports_weekly.html", user, db, lw=lw, qtd=qtd, dist_lw=dist_lw, dist_q=dist_q, by_lw=by_lw, by_q=by_q, sites=sites,
                   brand_of=brand_of, w_lw=w_lw, w_q=w_q, themes_used=themes_used, as_of=as_of or lw.end_date.isoformat(),
                   brands_f=brands_f, grp_f=grp_f, loc_f=loc_f, fq=fq, **_filter_ctx(db))


@app.get("/reports/alerts/preview", response_class=HTMLResponse)
def alerts_preview(user: User = Depends(auth.current_user), db: Session = Depends(get_db)):
    """What the instant-alert email looks like, built from whatever is waiting right now (or the
    most recent negatives when nothing is waiting). Nothing is sent or stamped."""
    from .alerts import AlertBatch
    b = collect_alerts(db)
    if b.empty:
        b = AlertBatch(negatives=db.execute(select(Review).join(ReviewSourceLink).where(ReviewSourceLink.active.is_(True), Review.is_deleted.is_(False),
                                                                                           Review.rating <= settings.negative_rating_max)
                                            .options(selectinload(Review.source_link).selectinload(ReviewSourceLink.location))
                                            .order_by(Review.created_at_source.desc()).limit(2)).scalars().all(),
                       removed=db.execute(select(Review).join(ReviewSourceLink).where(Review.is_deleted.is_(True))
                                          .options(selectinload(Review.source_link).selectinload(ReviewSourceLink.location))
                                          .order_by(Review.removed_at.desc()).limit(1)).scalars().all())
    return HTMLResponse(render_alert_html(b))


# ----------------------------------------------------------------- exports
EXPORT_COLUMNS = ["Review id", "Source", "Brand", "Site", "Rating", "Previous rating", "Posted", "Reviewer", "Review", "Our reply", "Replied at",
                  "Response hours", "Replied by", "Theme", "Employees mentioned", "Status", "Edited", "Removed", "Reported", "Link"]


def _export_rows(db: Session, user: User, view: str, brands_f, scope_ids, rat_f, q: str, dr: Optional[DateRange]) -> List[List]:
    overdue_cut = datetime.utcnow() - timedelta(hours=settings.overdue_hours)
    base, order = _inbox_query(user, view if view in VIEW_KEYS else "all", brands_f, scope_ids, rat_f, q, dr, overdue_cut)
    rows = db.execute(base.order_by(*order).limit(20000)).scalars().all()
    out = []
    for r in rows:
        last = [x for x in r.responses if x.status == "posted"]
        by = last[-1].created_by.name if last and last[-1].created_by else ("outside app" if r.has_owner_reply else "")
        status = "removed" if r.is_deleted else "archived" if r.is_archived else "replied" if r.has_owner_reply else "open"
        out.append([r.external_id, r.source, r.location.brand if r.location else "", r.location.name if r.location else (r.source_link.display_name or ""),
                    r.rating or "", r.prev_rating or "", _local(r.created_at_source, "%Y-%m-%d %H:%M"), r.author_name or "", r.text or "",
                    r.owner_reply_text or "", _local(r.replied_at, "%Y-%m-%d %H:%M") if r.has_owner_reply else "",
                    round(r.response_hours, 1) if r.response_hours is not None else "", by, r.category or "", "; ".join(r.mention_names), status,
                    "yes" if r.was_edited else "", _local(r.removed_at, "%Y-%m-%d") if r.removed_at else "", r.report_status or "",
                    f"{settings.app_base_url}/reviews/{r.id}"])
    return out


def _export_filters(db, brand, location_id, group_id, rating, range, start, end):
    brands_f, loc_f, grp_f, rat_f = _strs(brand), _ints(location_id), _ints(group_id), _ints(rating)
    dr = _dr(range, start, end) if (range or start or end) else None
    return brands_f, _scope(db, grp_f, loc_f), rat_f, dr


def _csv_response(name: str, header: List[str], rows: List[List]) -> RawResponse:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(rows)
    return RawResponse(buf.getvalue(), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": f'attachment; filename="{name}"'})


def _xlsx_response(name: str, sheets: List) -> RawResponse:
    """sheets = [(title, header, rows)]"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for title, header, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        for i, h in enumerate(header, start=1):
            width = max(len(str(h)), *(min(60, len(str(r[i - 1]))) for r in rows[:500])) if rows else len(str(h))
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(8, width + 2))
    out = io.BytesIO()
    wb.save(out)
    return RawResponse(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.get("/export/reviews.{fmt}")
def export_reviews(fmt: str, user: User = Depends(auth.current_user), db: Session = Depends(get_db), view: str = "all",
                   brand: List[str] = Query([]), location_id: List[str] = Query([]), group_id: List[str] = Query([]),
                   rating: List[str] = Query([]), q: str = "", range: str = "", start: str = "", end: str = ""):
    """Every review matching the inbox filters, one row each. CSV or Excel."""
    brands_f, scope_ids, rat_f, dr = _export_filters(db, brand, location_id, group_id, rating, range, start, end)
    rows = _export_rows(db, user, view, brands_f, scope_ids, rat_f, q, dr)
    stamp = datetime.now(_tz).strftime("%Y-%m-%d")
    if fmt == "xlsx":
        return _xlsx_response(f"reviews-{stamp}.xlsx", [("Reviews", EXPORT_COLUMNS, rows)])
    return _csv_response(f"reviews-{stamp}.csv", EXPORT_COLUMNS, rows)


@app.get("/export/distribution.{fmt}")
def export_distribution(fmt: str, user: User = Depends(auth.current_user), db: Session = Depends(get_db), brand: List[str] = Query([]),
                        location_id: List[str] = Query([]), group_id: List[str] = Query([]), range: str = "last30", start: str = "", end: str = ""):
    """Per-site star counts for the window, plus the negative-reason matrix. CSV or Excel."""
    brands_f, scope_ids, _rat, dr = _export_filters(db, brand, location_id, group_id, [], range, start, end)
    dr = dr or _dr("last30", "", "")
    dist = rating_distribution(db, dr, brands_f or None, scope_ids)
    header = ["Brand", "Site", "1 star", "2 star", "3 star", "4 star", "5 star", "Unrated", "Total", "Average", "3 star or below %", "Google shows avg", "Google shows count"]
    rows = [[r["brand"], r["name"], *r["counts"], r["unrated"], r["total"], r["avg"] or "", r["neg_pct"], r["listing_avg"] or "", r["listing_total"] or ""] for r in dist["rows"]]
    wr = window_report(db, dr, brands_f or None, scope_ids)
    theme_header = ["Site"] + list(THEMES)
    theme_rows = [[site] + [wr.theme_matrix.get(site, {}).get(c, 0) for c in THEMES] for site in sorted(wr.theme_matrix)]
    stamp = f"{dr.start_date.isoformat()}_{dr.end_date.isoformat()}"
    if fmt == "xlsx":
        return _xlsx_response(f"distribution-{stamp}.xlsx", [("By site", header, rows), ("Negative reasons", theme_header, theme_rows)])
    return _csv_response(f"distribution-{stamp}.csv", header, rows)


@app.get("/reports/morning", response_class=HTMLResponse)
def morning_preview(user: User = Depends(auth.current_user), db: Session = Depends(get_db), edition: str = "all", range: str = "yesterday",
                    start: str = "", end: str = ""):
    """Exactly what the email recipients see. `edition` = il / tn / all."""
    d = build_digest(db, edition if edition in EDITIONS else "all", _dr(range, start, end))
    return HTMLResponse(render_digest_html(d))




@app.get("/reports/employees", response_class=HTMLResponse)
def employees_report(request: Request, user: User = Depends(auth.current_user), db: Session = Depends(get_db), brand: List[str] = Query([]),
                     group_id: List[str] = Query([]), location_id: List[str] = Query([]), range: str = "last30", start: str = "", end: str = ""):
    dr = _dr(range, start, end)
    brands_f, grp_f, loc_f = _strs(brand), _ints(group_id), _ints(location_id)
    ids = _scope(db, grp_f, loc_f)
    rep = employee_report(db, dr, brands=brands_f or None, location_ids=ids)
    fq = "&".join([dr.query()] + [f"brand={b}" for b in brands_f] + [f"group_id={i}" for i in grp_f] + [f"location_id={i}" for i in loc_f])
    return _render(request, "reports_employees.html", user, db, rep=rep, dr=dr, brands_f=brands_f, grp_f=grp_f, loc_f=loc_f, fq=fq, **_filter_ctx(db))


@app.get("/sites/{location_id}", response_class=HTMLResponse)
def site_page(location_id: int, request: Request, user: User = Depends(auth.current_user), db: Session = Depends(get_db),
              range: str = "last30", start: str = "", end: str = "", unit: str = ""):
    loc = db.execute(select(Location).where(Location.id == location_id).options(selectinload(Location.sources))).scalar_one_or_none()
    if not loc:
        raise HTTPException(404, "Site not found")
    dr = _dr(range, start, end)
    wr = window_report(db, dr, location_ids=[loc.id])
    trends = build_trends(db, dr, location_ids=[loc.id], unit=unit or None)
    emp = employee_report(db, dr, location_ids=[loc.id])
    return _render(request, "site.html", user, db, loc=loc, w=wr, dr=dr, trends=trends, trends_json=json.dumps(trends), emp=emp,
                   unit=trends["unit"], fq=dr.query())


# ----------------------------------------------------------------- locations / sync admin
@app.get("/locations", response_class=HTMLResponse)
def locations(request: Request, user: User = Depends(auth.current_user), db: Session = Depends(get_db)):
    links = db.execute(select(ReviewSourceLink).options(selectinload(ReviewSourceLink.location))
                       .order_by(ReviewSourceLink.active.desc(), ReviewSourceLink.source, ReviewSourceLink.display_name)).scalars().all()
    locs = db.execute(select(Location).order_by(Location.brand, Location.name)).scalars().all()
    runs = db.execute(select(SyncRun).order_by(SyncRun.started_at.desc()).limit(30)).scalars().all()
    return _render(request, "locations.html", user, db, links=links, locs=locs, runs=runs)


@app.post("/locations/{link_id}/link")
def link_location(link_id: int, location_id: int = Form(...), user: User = Depends(auth.admin_user), db: Session = Depends(get_db)):
    link = db.get(ReviewSourceLink, link_id)
    if not link:
        raise HTTPException(404)
    link.location_id = location_id or None
    title = (link.display_name or "").lower()
    excluded = any(pat.lower() in title for pat in settings.listing_exclude_patterns)
    link.active = bool(link.location_id) and not excluded
    db.commit()
    return RedirectResponse(url="/locations", status_code=303)


def _background_sync(full: bool) -> None:
    with session_scope() as s:
        sync_all(s, full=full)


@app.post("/sync")
def trigger_sync(background: BackgroundTasks, full: int = Form(0), user: User = Depends(auth.admin_user)):
    background.add_task(_background_sync, bool(full))
    return RedirectResponse(url="/locations", status_code=303)


# ----------------------------------------------------------------- templates
@app.get("/admin/templates", response_class=HTMLResponse)
def admin_templates(request: Request, user: User = Depends(auth.current_user), db: Session = Depends(get_db), edit: int = 0):
    tpls = db.execute(select(ReplyTemplate).options(selectinload(ReplyTemplate.updated_by))
                      .order_by(ReplyTemplate.min_rating.desc(), ReplyTemplate.sort_order, ReplyTemplate.name)).scalars().all()
    groups: Dict[str, List[ReplyTemplate]] = {}
    for t in tpls:
        groups.setdefault(t.rating_band, []).append(t)
    ordered = [(b, groups[b]) for b in ("5 star", "4 star", "1-3 star", "Any rating") if b in groups]
    editing = db.get(ReplyTemplate, edit) if edit else None
    return _render(request, "admin_templates.html", user, db, groups=ordered, brands=_brands(db), editing=editing,
                   tag_options=TEMPLATE_TAGS, total=len(tpls))


@app.post("/admin/templates")
def save_template(user: User = Depends(auth.admin_user), db: Session = Depends(get_db), id: int = Form(0),
                  name: str = Form(...), brand: str = Form(""), min_rating: int = Form(1), max_rating: int = Form(5),
                  body: str = Form(...), sort_order: int = Form(100), active: int = Form(1), tags: List[str] = Form([])):
    t = db.get(ReplyTemplate, id) if id else ReplyTemplate()
    if not id:
        db.add(t)
    t.name, t.brand, t.body = name.strip(), (brand or None), body.strip()
    t.min_rating, t.max_rating = min(min_rating, max_rating), max(min_rating, max_rating)
    t.sort_order, t.active = sort_order, bool(active)
    t.tags = ";".join(tags) or None
    t.updated_by_id = user.id
    db.commit()
    return RedirectResponse(url="/admin/templates?msg=Template+saved", status_code=303)


@app.post("/admin/templates/{tid}/delete")
def delete_template(tid: int, user: User = Depends(auth.admin_user), db: Session = Depends(get_db)):
    t = db.get(ReplyTemplate, tid)
    if t:
        db.delete(t)
        db.commit()
    return RedirectResponse(url="/admin/templates", status_code=303)


# ----------------------------------------------------------------- admin: users / recipients / groups / employees / ai
@app.get("/admin/users", response_class=HTMLResponse)
def admin_users(request: Request, user: User = Depends(auth.admin_user), db: Session = Depends(get_db), msg: str = "", link: str = ""):
    users = db.execute(select(User).order_by(User.name)).scalars().all()
    return _render(request, "admin_users.html", user, db, users=users, msg=msg, link=link if link.startswith(settings.app_base_url) else "",
                   sso_enabled=settings.sso_enabled, domains=settings.sso_allowed_domains, mail_enabled=settings.mail_enabled)


@app.post("/admin/users")
def save_user(user: User = Depends(auth.admin_user), db: Session = Depends(get_db), email: str = Form(...),
              name: str = Form(...), password: str = Form(""), role: str = Form("agent")):
    """Create or update a user. A new user gets a welcome email with a set-password link; when
    email is not configured the link is shown to the admin to pass on."""
    from .account_mail import send_welcome
    email = email.strip().lower()
    role = role if role in ("agent", "admin") else "agent"
    u = db.execute(select(User).where(User.email == email)).scalar_one_or_none()
    if u is None:
        if password and auth.password_problem(password):
            return RedirectResponse(url=f"/admin/users?{urlencode({'msg': auth.password_problem(password)})}", status_code=303)
        sso_only = not password and settings.sso_enabled and auth.email_domain_allowed(email)
        u = User(email=email, name=name.strip(), password_hash=auth.hash_password(password) if password else None, role=role,
                 auth_provider="microsoft" if sso_only else "local")
        db.add(u)
        db.commit()
        link = auth.password_link(u, "welcome")
        if send_welcome(u, link, by=user.name):
            q = {"msg": f"Welcome email sent to {u.email} with a link to set their password."}
        else:
            q = {"msg": f"{u.name} added.", "link": link}
        return RedirectResponse(url=f"/admin/users?{urlencode(q)}", status_code=303)
    u.name, u.role = name.strip(), role
    if password:
        problem = auth.password_problem(password)
        if problem:
            return RedirectResponse(url=f"/admin/users?{urlencode({'msg': problem})}", status_code=303)
        u.password_hash = auth.hash_password(password)
    db.commit()
    return RedirectResponse(url="/admin/users?msg=Saved", status_code=303)


@app.post("/admin/users/{uid}/reset-link")
def admin_reset_link(uid: int, user: User = Depends(auth.admin_user), db: Session = Depends(get_db)):
    """Email (or show) a set-password link for a user who is locked out or new."""
    from .account_mail import send_reset
    u = db.get(User, uid)
    if not u or not u.active:
        return RedirectResponse(url="/admin/users?msg=User+not+found", status_code=303)
    link = auth.password_link(u, "welcome" if not u.password_hash else "reset")
    if send_reset(u, link):
        q = {"msg": f"Reset link emailed to {u.email}."}
    else:
        q = {"msg": f"Reset link for {u.email}:", "link": link}
    return RedirectResponse(url=f"/admin/users?{urlencode(q)}", status_code=303)


@app.post("/admin/users/{uid}/toggle")
def toggle_user(uid: int, user: User = Depends(auth.admin_user), db: Session = Depends(get_db)):
    u = db.get(User, uid)
    if u and u.id != user.id:
        u.active = not u.active
        db.commit()
    return RedirectResponse(url="/admin/users", status_code=303)


EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-']+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")


def _parse_recipients(raw: str):
    """'Jane <jane@x.com>, sam@y.com' (commas, semicolons or newlines) -> [(email, name)], [tokens with no address]."""
    found, bad = [], []
    for chunk in re.split(r"[,;\n]+", raw or ""):
        chunk = chunk.strip()
        if not chunk:
            continue
        hits = list(EMAIL_RE.finditer(chunk))
        if not hits:
            bad.append(chunk[:40])
        elif len(hits) == 1:
            found.append((hits[0].group(0).lower(), chunk[:hits[0].start()].strip(" <>\"'")))
        else:
            found.extend((m.group(0).lower(), "") for m in hits)
    return found, bad


def _legacy_brands(eds: List[str]) -> Optional[str]:
    """Keep the old `brands` column coherent with the editions (nothing reads it any more)."""
    if any(EDITIONS[e]["brands"] is None for e in eds if e in EDITIONS):
        return None
    return ";".join(sorted({b for e in eds if e in EDITIONS for b in EDITIONS[e]["brands"]})) or None


@app.get("/admin/recipients", response_class=HTMLResponse)
def admin_recipients(request: Request, user: User = Depends(auth.admin_user), db: Session = Depends(get_db),
                     added: int = 0, existing: int = 0, skipped: str = "", ed: str = ""):
    rows = db.execute(select(ReportRecipient).order_by(ReportRecipient.email)).scalars().all()
    by_edition = {key: [r for r in rows if key in r.editions and r.active] for key in RECIPIENT_LISTS}
    site_counts = Counter(l.brand for l in db.execute(select(Location)).scalars().all())
    desc = {key: (", ".join(f"{b} · {site_counts.get(b, 0)} sites" for b in e["brands"]) if e["brands"]
                  else f"All {sum(site_counts.values())} sites, grouped by brand") for key, e in EDITIONS.items()}
    desc[ALERTS_KEY] = f"New 1–{settings.negative_rating_max}★ reviews, reviews that vanish from Google, listings failing to sync. Sent minutes after the sync that notices."
    notice = {"added": added, "existing": existing, "skipped": [t for t in skipped.split("|") if t],
              "ed": ed if ed in RECIPIENT_LISTS else "all"} if (added or existing or skipped) else None
    return _render(request, "admin_recipients.html", user, db, by_edition=by_edition, desc=desc, notice=notice,
                   tz_short=settings.timezone.split("/")[-1].replace("_", " "))


@app.post("/admin/recipients")
def save_recipients(user: User = Depends(auth.admin_user), db: Session = Depends(get_db), emails: str = Form(""),
                    email: str = Form(""), name: str = Form(""), edition: str = Form("all")):
    """Add one or many addresses to an edition (or the alerts list). One address may sit on several."""
    ed = edition if edition in RECIPIENT_LISTS else "all"
    found, bad = _parse_recipients("\n".join(x for x in (emails, email) if x))
    if name.strip() and len(found) == 1:
        found = [(found[0][0], name.strip())]
    added = existing = 0
    for addr, nm in found:
        r = db.execute(select(ReportRecipient).where(ReportRecipient.email == addr)).scalar_one_or_none()
        if r is None:
            r = ReportRecipient(email=addr, edition="")
            db.add(r)
        eds = r.editions
        if ed in eds and r.active:
            existing += 1
        else:
            added += 1
            eds.append(ed)
        r.set_editions(eds)
        r.brands = _legacy_brands(r.editions)
        r.name = nm or r.name
        r.active = True
    db.commit()
    q = urlencode({"added": added, "existing": existing, "skipped": "|".join(bad), "ed": ed})
    return RedirectResponse(url=f"/admin/recipients?{q}", status_code=303)


@app.post("/admin/recipients/{rid}/delete")
def delete_recipient(rid: int, user: User = Depends(auth.admin_user), db: Session = Depends(get_db), edition: str = Form("")):
    """Drop someone from one edition, or from the list entirely when no edition is given."""
    r = db.get(ReportRecipient, rid)
    if r:
        eds = [e for e in r.editions if e != edition] if edition else []
        if eds:
            r.set_editions(eds)
            r.brands = _legacy_brands(eds)
        else:
            db.delete(r)
        db.commit()
    return RedirectResponse(url="/admin/recipients", status_code=303)


@app.get("/admin/groups", response_class=HTMLResponse)
def admin_groups(request: Request, user: User = Depends(auth.admin_user), db: Session = Depends(get_db), edit: int = 0):
    groups = db.execute(select(SiteGroup).options(selectinload(SiteGroup.locations)).order_by(SiteGroup.name)).scalars().all()
    locs = db.execute(select(Location).where(Location.active.is_(True)).order_by(Location.brand, Location.name)).scalars().all()
    editing = db.get(SiteGroup, edit) if edit else None
    return _render(request, "admin_groups.html", user, db, groups=groups, locs=locs, editing=editing)


@app.post("/admin/groups")
def save_group(user: User = Depends(auth.admin_user), db: Session = Depends(get_db), id: int = Form(0), name: str = Form(...),
               description: str = Form(""), location_ids: List[int] = Form([])):
    g = db.get(SiteGroup, id) if id else SiteGroup()
    if not id:
        db.add(g)
    g.name, g.description = name.strip()[:80], description.strip() or None
    g.locations = db.execute(select(Location).where(Location.id.in_(location_ids or [-1]))).scalars().all()
    db.commit()
    return RedirectResponse(url="/admin/groups", status_code=303)


@app.post("/admin/groups/{gid}/delete")
def delete_group(gid: int, user: User = Depends(auth.admin_user), db: Session = Depends(get_db)):
    g = db.get(SiteGroup, gid)
    if g:
        db.delete(g)
        db.commit()
    return RedirectResponse(url="/admin/groups", status_code=303)


@app.get("/admin/employees", response_class=HTMLResponse)
def admin_employees(request: Request, user: User = Depends(auth.admin_user), db: Session = Depends(get_db), edit: int = 0):
    employees = db.execute(select(Employee).options(selectinload(Employee.location)).order_by(Employee.active.desc(), Employee.name)).scalars().all()
    locs = db.execute(select(Location).where(Location.active.is_(True)).order_by(Location.brand, Location.name)).scalars().all()
    since = datetime.utcnow() - timedelta(days=180)
    cand_rows = db.execute(
        select(ReviewMention.name, func.count(ReviewMention.id), func.max(Review.created_at_source))
        .join(Review, ReviewMention.review_id == Review.id)
        .where(ReviewMention.employee_id.is_(None), Review.created_at_source >= since)
        .group_by(ReviewMention.name).order_by(func.count(ReviewMention.id).desc()).limit(40)).all()
    # most common site per candidate
    candidates = []
    for name, n, last in cand_rows:
        site = db.execute(select(Location.id, Location.name, func.count(Review.id))
                          .join(ReviewSourceLink, ReviewSourceLink.location_id == Location.id)
                          .join(Review, Review.source_link_id == ReviewSourceLink.id)
                          .join(ReviewMention, ReviewMention.review_id == Review.id)
                          .where(ReviewMention.name == name).group_by(Location.id, Location.name)
                          .order_by(func.count(Review.id).desc()).limit(1)).first()
        candidates.append({"name": name, "count": n, "last": last, "location_id": site[0] if site else None, "site": site[1] if site else ""})
    editing = db.get(Employee, edit) if edit else None
    return _render(request, "admin_employees.html", user, db, employees=employees, locs=locs, candidates=candidates, editing=editing)


@app.post("/admin/employees")
def save_employee(user: User = Depends(auth.admin_user), db: Session = Depends(get_db), id: int = Form(0), name: str = Form(...),
                  aliases: str = Form(""), location_id: int = Form(0), role: str = Form(""), active: int = Form(1)):
    e = db.get(Employee, id) if id else Employee()
    if not id:
        db.add(e)
    e.name, e.aliases, e.role = name.strip()[:80], aliases.strip() or None, role.strip() or None
    e.location_id, e.active = location_id or None, bool(active)
    db.flush()
    # link existing mentions that match any of the names
    names = {n.lower() for n in e.all_names()}
    for m in db.execute(select(ReviewMention).where(ReviewMention.employee_id.is_(None))).scalars().all():
        if m.name.lower() in names:
            m.employee_id = e.id
            m.name = e.name
    db.commit()
    return RedirectResponse(url="/admin/employees", status_code=303)


@app.post("/admin/employees/{eid}/delete")
def delete_employee(eid: int, user: User = Depends(auth.admin_user), db: Session = Depends(get_db)):
    e = db.get(Employee, eid)
    if e:
        for m in db.execute(select(ReviewMention).where(ReviewMention.employee_id == e.id)).scalars().all():
            m.employee_id = None
        db.delete(e)
        db.commit()
    return RedirectResponse(url="/admin/employees", status_code=303)


@app.post("/admin/employees/reindex")
def reindex_mentions(background: BackgroundTasks, user: User = Depends(auth.admin_user)):
    def _run():
        with session_scope() as s:
            from .text_intel import build_roster
            roster = build_roster(s)
            emps = {e.name: e.id for e in s.execute(select(Employee)).scalars().all()}
            for r in s.execute(select(Review).where(Review.is_deleted.is_(False)).options(selectinload(Review.mentions))).scalars().all():
                apply_intel(s, r, roster, emps)
    background.add_task(_run)
    return RedirectResponse(url="/admin/employees", status_code=303)


@app.get("/admin/ai", response_class=HTMLResponse)
def admin_ai(request: Request, user: User = Depends(auth.admin_user), db: Session = Depends(get_db)):
    rules = db.execute(select(AiRule).order_by(AiRule.sort_order, AiRule.id)).scalars().all()
    stats = db.execute(select(func.count(ReplyRow.id)).where(ReplyRow.ai_generated.is_(True), ReplyRow.status == "posted")).scalar() or 0
    return _render(request, "admin_ai.html", user, db, rules=rules, ai_enabled=settings.ai_enabled, model=settings.ai_model,
                   phones=settings.brand_phones, ai_posted=stats)


@app.post("/admin/ai/classify")
def ai_classify_window(background: BackgroundTasks, user: User = Depends(auth.admin_user), range: str = Form("last30"),
                       start: str = Form(""), end: str = Form(""), force: int = Form(0), back: str = Form("/reports")):
    if not settings.ai_enabled:
        raise HTTPException(400, "AI is not configured")
    dr = _dr(range, start, end)

    def _run():
        from .ai import AiUnavailable, classify_negative
        with session_scope() as s:
            q = select(Review).where(Review.is_deleted.is_(False), Review.rating <= settings.negative_rating_max,
                                     Review.created_at_source >= dr.start, Review.created_at_source < dr.end)
            for r in s.execute(q).scalars().all():
                if not (r.text or "").strip():
                    r.category = "No Content"; continue
                if r.category and not force and r.category != "Unknown":
                    continue
                try:
                    r.category = classify_negative(r, THEMES)
                except AiUnavailable:
                    break
    background.add_task(_run)
    return RedirectResponse(url=back, status_code=303)


@app.post("/admin/ai")
def save_rule(user: User = Depends(auth.admin_user), db: Session = Depends(get_db), id: int = Form(0), text: str = Form(...),
              sort_order: int = Form(100), active: int = Form(1)):
    r = db.get(AiRule, id) if id else AiRule()
    if not id:
        db.add(r)
    r.text, r.sort_order, r.active = text.strip(), sort_order, bool(active)
    db.commit()
    return RedirectResponse(url="/admin/ai", status_code=303)


@app.post("/admin/ai/{rid}/delete")
def delete_rule(rid: int, user: User = Depends(auth.admin_user), db: Session = Depends(get_db)):
    r = db.get(AiRule, rid)
    if r:
        db.delete(r)
        db.commit()
    return RedirectResponse(url="/admin/ai", status_code=303)


# ----------------------------------------------------------------- admin: sites & listings
@app.get("/admin/sites", response_class=HTMLResponse)
def admin_sites(request: Request, user: User = Depends(auth.admin_user), db: Session = Depends(get_db), edit: int = 0):
    locs = db.execute(select(Location).options(selectinload(Location.sources)).order_by(Location.active.desc(), Location.brand, Location.name)).scalars().all()
    unmapped = db.execute(select(ReviewSourceLink).where(ReviewSourceLink.location_id.is_(None)).order_by(ReviewSourceLink.source, ReviewSourceLink.display_name)).scalars().all()
    return _render(request, "admin_sites.html", user, db, locs=locs, unmapped=unmapped, editing=db.get(Location, edit) if edit else None,
                   brands=_brands(db), google_ready=True, facebook_ready=bool(settings.facebook_access_token),
                   excluded_patterns=settings.listing_exclude_patterns)


@app.post("/admin/sites")
def save_site(user: User = Depends(auth.admin_user), db: Session = Depends(get_db), id: int = Form(0), name: str = Form(...), brand: str = Form(...),
              state: str = Form(...), city: str = Form(""), snowflake_location_ids: str = Form(""), active: int = Form(1)):
    name = name.strip()[:120]
    loc = db.get(Location, id) if id else None
    if loc is None:
        dup = db.execute(select(Location).where(Location.name == name)).scalar_one_or_none()
        if dup:
            return RedirectResponse(url=f"/admin/sites?{urlencode({'msg': f'A site named {name} already exists.'})}", status_code=303)
        loc = Location(name=name)
        db.add(loc)
    loc.name, loc.brand, loc.state = name, brand.strip()[:40], state.strip().upper()[:2]
    loc.city = city.strip() or None
    loc.snowflake_location_ids = ";".join(x.strip() for x in snowflake_location_ids.replace(",", ";").split(";") if x.strip()) or None
    loc.active = bool(active)
    db.commit()
    return RedirectResponse(url=f"/admin/sites?{urlencode({'msg': f'Saved {loc.name}.'})}", status_code=303)


@app.post("/admin/sites/{sid}/toggle")
def toggle_site(sid: int, user: User = Depends(auth.admin_user), db: Session = Depends(get_db)):
    loc = db.get(Location, sid)
    if loc:
        loc.active = not loc.active
        for l in loc.sources:
            l.active = loc.active and not any(pat.lower() in (l.display_name or "").lower() for pat in settings.listing_exclude_patterns)
        db.commit()
    return RedirectResponse(url="/admin/sites", status_code=303)


@app.post("/admin/sites/discover")
def discover_listings(user: User = Depends(auth.admin_user), db: Session = Depends(get_db), source: str = Form("google")):
    """Ask the platform for every listing this account manages and record them. Runs now, a few seconds."""
    from . import discovery
    try:
        totals = discovery.discover_google(db) if source == "google" else discovery.discover_facebook(db)
        db.commit()
        msg = (f"Google: {totals['listings']} listings across {totals['accounts']} account(s), {totals['mapped']} mapped to sites." if source == "google"
               else f"Facebook: {totals['pages']} pages, {totals['mapped']} mapped to sites.")
    except Exception as exc:
        db.rollback()
        log.exception("discovery failed")
        msg = f"{source.capitalize()} discovery failed: {str(exc)[:200]}"
    return RedirectResponse(url=f"/admin/sites?{urlencode({'msg': msg})}", status_code=303)


@app.post("/admin/sites/listing")
def add_listing(user: User = Depends(auth.admin_user), db: Session = Depends(get_db), source: str = Form("google"), external_location_id: str = Form(...),
                external_account_id: str = Form(""), display_name: str = Form(""), listing_url: str = Form(""), location_id: int = Form(0)):
    """Record a listing by id when discovery cannot see it (e.g. a profile owned by another account)."""
    from . import discovery
    ext = external_location_id.strip().split("/")[-1]
    loc = db.get(Location, location_id) if location_id else None
    link = discovery.upsert_link(db, source.strip().lower(), ext, display_name.strip() or ext, external_account_id=external_account_id.strip().split("/")[-1] or None,
                                 listing_url=listing_url.strip() or None, location=loc, auto_map=False)
    if loc and not discovery.excluded(link.display_name or ""):
        link.location_id, link.active = loc.id, True
    db.commit()
    return RedirectResponse(url=f"/admin/sites?{urlencode({'msg': f'Listing {link.display_name} recorded' + (f' and mapped to {loc.name}.' if loc else '; map it to a site below.')})}", status_code=303)


# ----------------------------------------------------------------- admin: API keys
@app.get("/admin/api", response_class=HTMLResponse)
def admin_api(request: Request, user: User = Depends(auth.admin_user), db: Session = Depends(get_db), new: str = ""):
    keys = db.execute(select(ApiKey).options(selectinload(ApiKey.created_by)).order_by(ApiKey.active.desc(), ApiKey.created_at.desc())).scalars().all()
    return _render(request, "admin_api.html", user, db, keys=keys, new_key=new, base=settings.app_base_url, cors=settings.api_cors_origins)


@app.post("/admin/api")
def create_api_key(request: Request, user: User = Depends(auth.admin_user), db: Session = Depends(get_db), name: str = Form(...)):
    raw = _new_api_key()
    db.add(ApiKey(name=name.strip()[:120] or "Unnamed", prefix=raw[:11], key_hash=_hash_api_key(raw), created_by_id=user.id))
    db.commit()
    request.session_new_key = raw  # type: ignore[attr-defined]
    return RedirectResponse(url=f"/admin/api?{urlencode({'new': raw})}", status_code=303)


@app.post("/admin/api/{kid}/revoke")
def revoke_api_key(kid: int, user: User = Depends(auth.admin_user), db: Session = Depends(get_db)):
    k = db.get(ApiKey, kid)
    if k:
        k.active = False
        db.commit()
    return RedirectResponse(url="/admin/api?msg=Key+revoked", status_code=303)
